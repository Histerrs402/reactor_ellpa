"""Чтение книги Excel только на чтение. Исходный файл не изменяется: программа
копирует его во временную папку и работает с копией."""
from __future__ import annotations

import hashlib
import re
import shutil
import uuid
import warnings
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException


class UserError(Exception):
    """Ошибка, понятная пользователю (показывается в интерфейсе как есть)."""


SUPPORTED = (".xlsx", ".xlsm")


def stage_copy(path: str, tmp_dir: Path) -> tuple[Path, str, int, datetime]:
    src = Path(path.strip().strip('"')) if path else None
    if not src or not str(src):
        raise UserError("Не указан файл Excel.")
    if not src.exists():
        raise UserError(f"Файл не найден: {src}")
    if src.is_dir():
        raise UserError(f"Указана папка, а не файл: {src}")
    if src.suffix.lower() not in SUPPORTED:
        raise UserError(
            f"Неподдерживаемый формат «{src.suffix}». Нужен файл .xlsx или .xlsm "
            "(старый .xls сохраните в Excel как .xlsx)."
        )
    tmp_dir.mkdir(parents=True, exist_ok=True)
    dst = tmp_dir / f"import_{uuid.uuid4().hex}{src.suffix.lower()}"
    try:
        shutil.copyfile(src, dst)
    except PermissionError:
        raise UserError(
            "Нет доступа к файлу. Возможно, он открыт в Excel в монопольном режиме "
            "или сохраняется. Закройте его или подождите и повторите."
        )
    except OSError as e:
        raise UserError(f"Не удалось прочитать файл: {e}")
    h = hashlib.sha256()
    with open(dst, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    st = src.stat()
    return dst, h.hexdigest(), st.st_size, datetime.fromtimestamp(st.st_mtime).replace(microsecond=0)


def core_modified(path: Path) -> datetime | None:
    """Дата последнего сохранения книги из её свойств (docProps/core.xml), в местном времени."""
    try:
        with zipfile.ZipFile(path) as z:
            xml = z.read("docProps/core.xml").decode("utf-8", "replace")
    except (KeyError, zipfile.BadZipFile, OSError):
        return None
    m = re.search(r"<dcterms:modified[^>]*>([^<]+)</dcterms:modified>", xml)
    if not m:
        return None
    s = m.group(1).strip()
    try:
        dt = datetime.fromisoformat(s.replace("Z", "+00:00"))
    except ValueError:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone().replace(tzinfo=None, microsecond=0)


def _clean_header(v) -> str:
    if v is None:
        return ""
    return str(v).replace(" ", " ").strip()


def _key(s: str) -> str:
    return re.sub(r"[\s_]+", "", s).lower()


def resolve_column(headers: list[str], wanted: str | None) -> str | None:
    """Находит столбец по имени: точное совпадение, затем без учёта регистра и пробелов."""
    if not wanted:
        return None
    if wanted in headers:
        return wanted
    wk = _key(wanted)
    for h in headers:
        if h and _key(h) == wk:
            return h
    return None


class Workbook:
    def __init__(self, path: Path):
        self.path = path
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                self.wb = openpyxl.load_workbook(
                    str(path), read_only=True, data_only=True, keep_links=False
                )
        except (InvalidFileException, zipfile.BadZipFile, KeyError) as e:
            raise UserError(f"Файл не похож на книгу Excel или повреждён ({e.__class__.__name__}).")
        except Exception as e:  # noqa: BLE001
            raise UserError(f"Не удалось открыть книгу Excel: {e}")

    @property
    def sheet_names(self) -> list[str]:
        return list(self.wb.sheetnames)

    def sheet_state(self, name: str) -> str:
        try:
            return self.wb[name].sheet_state
        except Exception:  # noqa: BLE001
            return "visible"

    def headers(self, name: str) -> list[str]:
        ws = self.wb[name]
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
                hdr = [_clean_header(v) for v in row]
                while hdr and not hdr[-1]:
                    hdr.pop()
                return hdr
        return []

    def approx_rows(self, name: str) -> int | None:
        try:
            mr = self.wb[name].max_row
            return max(mr - 1, 0) if mr else None
        except Exception:  # noqa: BLE001
            return None

    def iter_rows(self, name: str, columns: list[str] | None = None, progress=None):
        """Строки листа как словари {заголовок: значение}. Полностью пустые строки пропускаются.
        columns — ограничить набором столбцов (ускоряет чтение больших листов)."""
        ws = self.wb[name]
        hdr = self.headers(name)
        if columns is None:
            idx = [(i, h) for i, h in enumerate(hdr) if h]
        else:
            want = set(columns)
            idx = [(i, h) for i, h in enumerate(hdr) if h in want]
        total = self.approx_rows(name)
        n = 0
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            for row in ws.iter_rows(min_row=2, values_only=True):
                n += 1
                if progress and n % 1000 == 0:
                    progress(n, total)
                if not row or all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
                    continue
                ln = len(row)
                yield n + 1, {h: (row[i] if i < ln else None) for i, h in idx}
        if progress:
            progress(n, total)

    def close(self):
        try:
            self.wb.close()
        except Exception:  # noqa: BLE001
            pass
