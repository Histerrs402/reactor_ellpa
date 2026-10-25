"""Генератор тестовых книг Excel со структурой, как у реальной выгрузки OData
(подмножество столбцов + персональные поля, чтобы проверить их исключение)."""
from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook

MAIN_HEADERS = [
    "ДатаСоздания", "УчетныйНомерРемонтаСистемный", "Бренд", "Модель", "ДатаПриемаВАсц",
    "ДатаЗавершенияРемонта", "ДатаГотовностиКВыдаче", "АсцВыполняющийРемонт", "КодАсцNsoВыполняющийРемонт",
    "ТипРемонта", "ТипРемонтаПриЗавершении", "СтатусРемонта", "КатегорияТехники", "ДатаСозданияЗаявкиОнЛайн", "ДатаВыездДоставкаОнЛайн",
    "ДатаДиагностикиОнЛайн", "ДатаРемонтаОнЛайн", "ДатаЗавершенияОнЛайн", "ДатаОтчетаОнЛайн", "ДатаЗаказаЗЧОнЛайн",
    "ДатаЗапросаНаОтменуОнЛайн", "НомерАкта", "ДатаАкта", "Страна", "Регион", "Город",
    "ФамилияКлиента", "ИмяКлиента", "МобТелефонКлиента", "АдресКлиента", "EMailКлиента",
    "ПовторныйРемонт", "ДатаИзменения", "TAT", "комментарий",
    "НомерАктаНрпАто", "ДатаСозданияАкта", "ПричинаАкта", "Деталь1ПартНомер", "СуммаЗЧ", "ПолитикаРемонта", "Id", "Транспортировка",
    "ДатаУтвержденияАкта",
]
STATUS_COL = {
    "Заявка": "ДатаСозданияЗаявкиОнЛайн",
    "Выезд/Доставка": "ДатаВыездДоставкаОнЛайн",
    "Диагностика": "ДатаДиагностикиОнЛайн",
    "Заказаны запчасти": "ДатаЗаказаЗЧОнЛайн",
    "Ремонт": "ДатаРемонтаОнЛайн",
    "Ремонт завершён": "ДатаЗавершенияОнЛайн",
    "Отчет": "ДатаОтчетаОнЛайн",
}


def repair(et: str, status: str, created: datetime, status_date: datetime | None, asc="АСЦ Альфа",
           model="ECAM22.110.B", **extra) -> dict:
    r = {
        "УчетныйНомерРемонтаСистемный": et,
        "СтатусРемонта": status,
        "ДатаСоздания": created,
        "АсцВыполняющийРемонт": asc,
        "КодАсцNsoВыполняющийРемонт": "A" + str(abs(hash(asc)) % 1000),
        "Модель": model,
        "Бренд": "DELONGHI",
        "ТипРемонта": "Клиентский гарантийный",
        "ТипРемонтаПриЗавершении": "Клиентский гарантийный",
        "КатегорияТехники": "CM FULLY-AUTOMATIC",
        "Страна": "РОССИЯ", "Регион": "Москва", "Город": "Москва",
        "ФамилияКлиента": "Иванов", "ИмяКлиента": "Иван", "МобТелефонКлиента": "+70000000000",
        "АдресКлиента": "ул. Тестовая, 1", "EMailКлиента": "test@example.com",
        "ПовторныйРемонт": 0,
        "ДатаСозданияЗаявкиОнЛайн": created,
    }
    if status_date is not None and status in STATUS_COL:
        r[STATUS_COL[status]] = status_date
    r["ДатаИзменения"] = status_date or created
    r.update(extra)
    return r


def make_workbook(path: Path, rows: list[dict], snapshot: datetime, comments: dict | None = None,
                  avr: list[dict] | None = None, with_tat: bool = True, headers: list[str] | None = None,
                  main_sheet: str = "DL_GR_open_reps") -> Path:
    hdr = list(headers or MAIN_HEADERS)
    if not with_tat and "TAT" in hdr:
        hdr.remove("TAT")
    wb = Workbook()
    ws = wb.active
    ws.title = main_sheet
    ws.append(hdr)
    for r in rows:
        r = dict(r)
        if with_tat and isinstance(r.get("ДатаСоздания"), datetime):
            r.setdefault("TAT", round((snapshot - r["ДатаСоздания"]).total_seconds() / 86400, 2))
        ws.append([r.get(h) for h in hdr])
    if comments is not None:
        cs = wb.create_sheet("Comments_DB")
        cs.append(["Repair", "Comment"])
        for et, text in comments.items():
            cs.append([et, text])
    if avr is not None:
        av = wb.create_sheet("DL_GR_AVR")
        av.append(["ДатаСоздания", "УчетныйНомерРемонтаСистемный", "СтатусРемонта", "НомерАкта", "ДатаАкта"])
        for a in avr:
            av.append([a.get("created"), a["et"], a.get("status", "Оплата"), a.get("act_number"), a.get("act_date")])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
