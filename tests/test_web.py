"""Проверка веб-интерфейса через тестовый клиент Flask: страницы, формы, фоновая загрузка."""
from __future__ import annotations

import io
import os
import re
import shutil
import sys
import tempfile
import time
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from test_core import S1, S2, rows_v1, rows_v2  # noqa: E402
from testdata import make_workbook  # noqa: E402


class TestWeb(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="rm_web_"))
        os.environ["REPAIR_MONITOR_DATA"] = str(self.tmp / "data")
        import app as A
        from monitor import jobs
        self.jobs = jobs
        jobs.JOB = jobs.ImportJob()
        self.app = A.create_app()
        self.app.config["TESTING"] = True
        self.c = self.app.test_client()
        self.f1 = make_workbook(self.tmp / "v1.xlsx", rows_v1(), S1, comments={"ET0000001": "из excel"},
                                avr=[])
        self.f2 = make_workbook(self.tmp / "v2.xlsx", rows_v2(), S2, comments={"ET0000001": "из excel"},
                                avr=[])

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def run_import(self, path, force=False):
        data = {"path": str(path)}
        if force:
            data["force"] = "1"
        r = self.c.post("/import/start", data=data)
        self.assertEqual(r.status_code, 302)
        for _ in range(200):
            st = self.c.get("/import/status").get_json()
            if not st["running"]:
                return st["result"]
            time.sleep(0.05)
        self.fail("загрузка не завершилась")

    def test_empty_state_and_pages(self):
        r = self.c.get("/")
        self.assertEqual(r.status_code, 200)
        self.assertIn("Начало работы", r.get_data(as_text=True))
        self.assertEqual(self.c.get("/repairs").status_code, 302)
        for u in ("/import", "/settings", "/history", "/comments"):
            self.assertEqual(self.c.get(u).status_code, 200, u)

    def test_check_import_and_navigation(self):
        r = self.c.post("/import/check", data={"path": str(self.f1)})
        html = r.get_data(as_text=True)
        self.assertIn("Структура подходит для загрузки", html)
        r = self.c.post("/import/check", data={"path": str(self.tmp / "nope.xlsx")})
        self.assertIn("Файл не найден", r.get_data(as_text=True))

        res = self.run_import(self.f1)
        self.assertEqual(res["status"], "ok", res)
        self.assertEqual(res["counts"]["initial"], 7)
        # файл запомнен для кнопки «Обновить данные»
        html = self.c.get("/").get_data(as_text=True)
        self.assertIn("Обновить данные", html)
        self.assertIn("Требуют внимания", html)
        for u in ("/", "/repairs", "/repairs?presence=all", "/attention", "/analytics", "/history",
                  "/history?initial=1", "/comments", "/repair/ET0000001", "/settings?tab=statuses",
                  "/settings?tab=mapping", "/settings?tab=source", "/repairs?q=0000004&basis=detected"):
            r = self.c.get(u)
            self.assertEqual(r.status_code, 200, u)
        self.assertEqual(self.c.get("/repair/ET9999999").status_code, 404)
        # поиск по ET
        html = self.c.get("/repairs?q=ET0000005").get_data(as_text=True)
        self.assertIn("ET0000005", html)
        self.assertNotIn("ET0000006", html)
        # повторная загрузка того же файла
        res = self.run_import(self.f1)
        self.assertEqual(res["status"], "duplicate")
        # вторая выгрузка
        res = self.run_import(self.f2)
        self.assertEqual(res["status"], "ok", res)
        self.assertEqual(res["counts"]["disappeared"], 2)
        html = self.c.get("/attention?reason=disappeared").get_data(as_text=True)
        self.assertIn("ET0000004", html)
        # ET0000003 был в статусе «Отчет» (закрыт) — его исчезновение ожидаемо, внимания не требует
        self.assertNotIn('href="/repair/ET0000003"', html)
        html = self.c.get("/history").get_data(as_text=True)
        self.assertIn("Смена статуса", html)
        self.assertIn("Смена АСЦ", html)

    def test_comments_and_resolve(self):
        self.run_import(self.f1)
        page = self.c.get("/repair/ET0000002").get_data(as_text=True)
        token = re.search(r'name="token" value="([0-9a-f]+)"', page).group(1)
        for _ in range(2):  # двойная отправка формы
            r = self.c.post("/repair/ET0000002/comment", data={"text": "Позвонила в АСЦ", "author": "Юлиана",
                                                                 "token": token})
            self.assertEqual(r.status_code, 302)
        html = self.c.get("/repair/ET0000002").get_data(as_text=True)
        self.assertEqual(html.count('<div class="text">Позвонила в АСЦ</div>'), 1)  # в журнале — одна запись
        self.assertIn("Комментарий в «Листе»: <b>Позвонила в АСЦ</b>", html)
        self.assertIn("Юлиана", html)
        # комментарий виден в реестре как последний
        self.assertIn("Позвонила в АСЦ", self.c.get("/repairs?q=ET0000002").get_data(as_text=True))
        # пустой комментарий не сохраняется; неизвестный ET — ошибка
        self.c.post("/repair/ET0000002/comment", data={"text": "  ", "token": "t2"})
        r = self.c.post("/comments", data={"et": "ET1234567", "text": "x", "token": "t3"}, follow_redirects=True)
        self.assertIn("не найден", r.get_data(as_text=True))
        r = self.c.post("/comments", data={"et": "et0000003", "text": "Жду акт", "token": "t4"}, follow_redirects=True)
        self.assertIn("Комментарий к ET0000003 сохранён", r.get_data(as_text=True))
        # комментарии переживают повторный импорт
        self.run_import(self.f2)
        html = self.c.get("/repair/ET0000002").get_data(as_text=True)
        self.assertIn("Позвонила в АСЦ", html)
        # отметить исчезнувший как проверенный и снять отметку
        self.assertIn("ET0000004", self.c.get("/attention").get_data(as_text=True))
        r = self.c.post("/repair/ET0000004/resolve", data={"action": "resolve", "note": "закрыт в системе",
                                                            "author": "Юлиана"})
        self.assertEqual(r.status_code, 302)
        att = self.c.get("/attention").get_data(as_text=True)
        self.assertNotIn('href="/repair/ET0000004"', att)
        card = self.c.get("/repair/ET0000004").get_data(as_text=True)
        self.assertIn("Отмечен как проверенный", card)
        self.c.post("/repair/ET0000004/resolve", data={"action": "reopen"})
        self.assertIn('href="/repair/ET0000004"', self.c.get("/attention").get_data(as_text=True))

    def test_settings_save_changes_calculation(self):
        self.run_import(self.f1)
        html = self.c.get("/attention?reason=stuck").get_data(as_text=True)
        self.assertIn('href="/repair/ET0000001"', html)
        # поднять общий порог до 10 — ET0000001 (9 дней) уходит
        r = self.c.post("/settings/stuck", data={"enabled": "1", "default_days": "10", "n_asc": "0", "n_sa": "0"})
        self.assertEqual(r.status_code, 302)
        html = self.c.get("/attention?reason=stuck").get_data(as_text=True)
        self.assertNotIn('href="/repair/ET0000001"', html)
        # некорректное число — настройки не сохраняются
        r = self.c.post("/settings/stuck", data={"enabled": "1", "default_days": "abc"}, follow_redirects=True)
        self.assertIn("не сохранены", r.get_data(as_text=True))
        # вкладка статусов: включить контроль «Отчет» с порогом 30
        page = self.c.get("/settings?tab=statuses").get_data(as_text=True)
        names = re.findall(r'name="st_(\d+)" value="([^"]+)"', page)
        form = {"n_status": str(len(names))}
        for idx, st in names:
            form[f"st_{idx}"] = st
            form[f"stuck_on_{idx}"] = "1"
            form[f"od_on_{idx}"] = "1" if st != "Отчет" else ""
            if st == "Отчет":
                form[f"stuck_days_{idx}"] = "30"
                form[f"sd_{idx}"] = "ДатаОтчетаОнЛайн"
        r = self.c.post("/settings/statuses", data=form)
        self.assertEqual(r.status_code, 302)
        html = self.c.get("/attention?reason=stuck").get_data(as_text=True)
        self.assertIn('href="/repair/ET0000003"', html)  # 91 день > 30
        # сопоставление: обязательное поле нельзя очистить
        r = self.c.post("/settings/mapping", data={"map_et": "", "map_status": "СтатусРемонта"}, follow_redirects=True)
        self.assertIn("не заданы обязательные поля", r.get_data(as_text=True))
        # резервная копия
        r = self.c.post("/backup", follow_redirects=True)
        self.assertIn("Резервная копия создана", r.get_data(as_text=True))

    def test_status_country_filters_and_demo(self):
        rows = rows_v1()
        rows[4]["Страна"] = "КАЗАХСТАН"  # ET0000005, статус «Ремонт»
        self.run_import(make_workbook(self.tmp / "c.xlsx", rows, S1, comments={"ET0000002": "Демо"}, avr=[]))

        def shown(url):
            html = self.c.get(url).get_data(as_text=True)
            return sorted(set(re.findall(r'href="/repair/(ET\d+)"', html)))
        everything = ["ET0000001", "ET0000002", "ET0000004", "ET0000005", "ET0000006", "ET0000007"]
        self.assertEqual(shown("/repairs"), everything)  # ET0000003 «Отчет» скрыт
        html = self.c.get("/repairs").get_data(as_text=True)
        self.assertNotIn('name="status" value="Отчет"', html)  # закрытый статус в срезе не предлагается
        self.assertIn('name="country" value="КАЗАХСТАН"', html)
        # несколько статусов; выбор запоминается
        self.assertEqual(shown("/repairs?status_set=1&status=Диагностика&status=Ремонт"),
                         ["ET0000001", "ET0000005", "ET0000006", "ET0000007"])
        self.assertEqual(shown("/repairs"), ["ET0000001", "ET0000005", "ET0000006", "ET0000007"])
        # страна сужает дальше, в том числе панель
        self.assertEqual(shown("/repairs?country_set=1&country=КАЗАХСТАН"), ["ET0000005"])
        dash = self.c.get("/").get_data(as_text=True)
        self.assertRegex(dash, r'Ремонтов в работе</div>\s*<div class="value">1</div>')
        self.assertIn("Сбросить всё", self.c.get("/repairs").get_data(as_text=True))
        # «Сбросить всё» снимает запомненные срезы
        self.assertEqual(shown("/repairs?rtype_set=1&status_set=1&country_set=1"), everything)
        # демо: без критичности, в «Требуют внимания» не попадает
        self.assertEqual(shown("/repairs?problem=demo"), ["ET0000002"])
        self.assertIn("демо", self.c.get("/repair/ET0000002").get_data(as_text=True))
        self.assertNotIn("ET0000002", shown("/attention"))
        # правило демо можно выключить в настройках
        r = self.c.post("/settings/general", data={"author": "", "backup_keep": "15", "demo_keywords": ""},
                        follow_redirects=True)
        self.assertIn("Настройки сохранены", r.get_data(as_text=True))
        self.assertIn("ET0000002", shown("/attention"))

    def test_sheet_comments_by_et(self):
        from openpyxl import load_workbook
        self.run_import(self.f1)
        html = self.c.get("/sheet").get_data(as_text=True)
        self.assertIn("ДатаСоздания", html)
        self.assertIn('data-et="ET0000002"', html)
        self.assertIn(">из excel</textarea>", html)  # комментарий из Comments_DB уже в строке
        self.assertNotIn('data-et="ET0000003"', html)  # «Отчет» — закрыт
        # комментарий в строке сохраняется по ET
        r = self.c.post("/api/comment", json={"et": "ET0000004", "text": "  написала в АСЦ 23/10 "})
        self.assertEqual(r.get_json()["text"], "написала в АСЦ 23/10")
        self.assertEqual(self.c.post("/api/comment", json={"et": "ET9999999", "text": "x"}).status_code, 404)
        html = self.c.get("/sheet?cm=with").get_data(as_text=True)
        self.assertIn('data-et="ET0000004"', html)
        self.assertIn("has-cm", html)
        # ремонт пропал из выгрузки — комментарий остался; в реестре и выгрузке комментариев он есть
        self.run_import(self.f2)
        html = self.c.get("/sheet?presence=all&cm=with").get_data(as_text=True)
        self.assertIn(">написала в АСЦ 23/10</textarea>", html)
        wb = load_workbook(io.BytesIO(self.c.get("/comments/export").data))
        self.assertEqual(wb.sheetnames, ["Comments_DB", "Комментарии", "История комментариев"])
        rows = list(wb["Comments_DB"].iter_rows(values_only=True))
        self.assertEqual(rows[0], ("Repair", "Comment"))
        self.assertIn(("ET0000004", "написала в АСЦ 23/10"), rows)
        info = {r[0]: r for r in wb["Комментарии"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(info["ET0000004"][8], "нет (исчез)")
        # правка и очистка: текущий меняется, история остаётся
        self.c.post("/api/comment", json={"et": "ET0000001", "text": "жду ответа"})
        self.c.post("/api/comment", json={"et": "ET0000001", "text": ""})
        wb = load_workbook(io.BytesIO(self.c.get("/comments/export").data))
        cur = {r[0]: r[1] for r in wb["Comments_DB"].iter_rows(min_row=2, values_only=True)}
        self.assertNotIn("ET0000001", cur)
        hist = [r[3] for r in wb["История комментариев"].iter_rows(min_row=2, values_only=True) if r[0] == "ET0000001"]
        self.assertIn("жду ответа", hist)
        self.assertTrue(any(h.startswith("Комментарий в листе очищен") for h in hist))
        # выгрузка листа: те же столбцы + комментарий, фильтр по столбцу учитывается
        self.c.post("/api/comment", json={"et": "ET0000002", "text": "жду ЗЧ"})
        wb = load_workbook(io.BytesIO(self.c.get("/sheet?export=xlsx&cf11=жду").data))
        ws = wb.active
        hdr = [c.value for c in ws[1]]
        self.assertEqual(hdr[:2], ["ДатаСоздания", "УчетныйНомерРемонтаСистемный"])
        self.assertEqual(hdr[-2:], ["комментарий", "комментарий изменён"])
        data = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual([(r[1], r[11]) for r in data], [("ET0000002", "жду ЗЧ")])
        self.assertIsInstance(data[0][0], __import__("datetime").datetime)
        # настройка столбцов листа
        r = self.c.post("/settings/sheet", data={"sheet_columns": "УчетныйНомерРемонтаСистемный\nСтатусРемонта\nTAT\nкомментарий"},
                        follow_redirects=True)
        self.assertIn("Настройки сохранены", r.get_data(as_text=True))
        html = self.c.get("/sheet").get_data(as_text=True)
        self.assertIn('title="СтатусРемонта"', html)
        self.assertNotIn('title="Модель"', html)

    def test_open_repairs_and_stuck(self):
        from datetime import datetime as D
        from openpyxl import load_workbook
        from testdata import repair
        rows = [
            repair("ET0000101", "Заявка", D(2026, 10, 1, 9), D(2026, 10, 1, 9), asc="АСЦ Альфа"),          # 18 дн. в заявке
            repair("ET0000102", "Заявка", D(2026, 10, 17, 9), D(2026, 10, 17, 9), asc="АСЦ Альфа"),        # 2 дн. — норма
            repair("ET0000103", "Диагностика", D(2026, 10, 2, 9), D(2026, 10, 5, 9), asc="АСЦ Бета",
                   ДатаПриемаВАсц=D(2026, 10, 4), Город="Казань"),                                          # принята 15 дн. назад
            repair("ET0000104", "Ремонт завершён", D(2026, 9, 1, 9), D(2026, 10, 10, 9), asc="АСЦ Бета",
                   ДатаПриемаВАсц=D(2026, 9, 3)),                                                            # 9 дн. без отчёта
            repair("ET0000105", "Отчет", D(2026, 8, 1, 9), D(2026, 9, 1, 9)),                               # закрыт
            repair("ET0000106", "Ремонт завершён", D(2026, 9, 1, 9), D(2026, 10, 1, 9), asc="АСЦ Бета"),  # демо
        ]
        self.run_import(make_workbook(self.tmp / "o.xlsx", rows, S1, comments={"ET0000106": "Демо"}, avr=[]))

        def shown(url):
            html = self.c.get(url).get_data(as_text=True)
            return sorted(set(re.findall(r'href="/repair/(ET\d+)"', html.split("<h2>Ремонты")[-1])))
        html = self.c.get("/open").get_data(as_text=True)
        self.assertNotIn("ET0000105", html)  # «Отчет» исключён
        order = [m for m in re.findall(r"<b>(Заявка|Диагностика|Ремонт завершён)</b>", html)]
        self.assertEqual(order, ["Заявка", "Диагностика", "Ремонт завершён"])  # порядок этапов
        self.assertIn("18", html)
        self.assertEqual(shown("/open?st=Заявка"), ["ET0000101", "ET0000102"])
        self.assertEqual(shown("/open?accepted=yes"), ["ET0000103", "ET0000104"])
        self.assertEqual(shown("/open?accepted=no&days_from=10"), ["ET0000101", "ET0000106"])
        self.assertEqual(shown("/open?city=Казань"), ["ET0000103"])
        wb = load_workbook(io.BytesIO(self.c.get("/open?sort=days&dir=desc&export=xlsx").data))
        data = list(wb.active.iter_rows(values_only=True))
        self.assertEqual(data[0][:5], ("ET", "Дата создания заявки", "Статус", "Дата перехода в статус", "Дней в статусе"))
        self.assertEqual(len(data), 6)
        # где застряли: порог по каждому статусу (по умолчанию 5 дн. и больше); демо и «Отчет» не учитываются
        html = self.c.get("/stuck").get_data(as_text=True)
        body = html.split("<h2>Ремонты")[-1]
        for et in ("ET0000101", "ET0000103", "ET0000104"):
            self.assertIn(et, body)
        for et in ("ET0000102", "ET0000105", "ET0000106"):
            self.assertNotIn(et, body)
        self.assertIn(">Принята АСЦ, нет диагностики<", body)
        # письмо в АСЦ: тема и текст
        html = self.c.get("/stuck?asc_sel=АСЦ Бета").get_data(as_text=True)
        self.assertIn('value="Запрос статуса ремонтов (2 шт.) — АСЦ Бета"', html)
        mail = html.split('class="mail-body">')[1].split("</textarea>")[0]
        self.assertIn("техника принята 04.10.2026, диагностики нет (15 дн.)", mail)
        self.assertIn("ET0000104, ECAM22.110.B: ремонт завершён 10.10.2026, не переведён в «Отчет» (9 дн.)", mail)
        self.assertIn("mailto:?subject=", html)
        # порог для статуса меняется на странице: «Ремонт завершён» — 10 дн. и больше
        self.c.post("/stuck/settings", data={"n_st": "1", "st_0": "Ремонт завершён", "on_0": "1", "days_0": "10",
                                              "acc_on": "1", "acc_days": "3", "skip_demo": "1",
                                              "acc_statuses": ["Заявка", "Выезд/Доставка", "Диагностика"]})
        self.assertNotIn("ET0000104", self.c.get("/stuck").get_data(as_text=True).split("<h2>Ремонты")[-1])
        # статус можно выключить
        self.c.post("/stuck/settings", data={"n_st": "1", "st_0": "Заявка", "days_0": "5", "acc_on": "1", "acc_days": "3",
                                              "skip_demo": "1"})
        self.assertNotIn("ET0000101", self.c.get("/stuck").get_data(as_text=True).split("<h2>Ремонты")[-1])
        wb = load_workbook(io.BytesIO(self.c.get("/stuck?export=xlsx").data))
        self.assertEqual(wb.sheetnames, ["Зависшие ремонты", "По АСЦ"])
        ets = sorted(r[2] for r in wb["Зависшие ремонты"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(ets, ["ET0000103", "ET0000103", "ET0000104"])
        # ремонт с комментарием — зелёным; фильтры под заголовками; пороги сворачиваются
        self.c.post("/api/comment", json={"et": "ET0000104", "text": "Ждём акт от АСЦ"})
        html = self.c.get("/stuck").get_data(as_text=True)
        self.assertIn('<details class="fold" id="thr-fold">', html)
        self.assertEqual(html.count('class="cf"'), 7)
        self.assertRegex(html, r'<tr class="has-cm">\s*<td data-f="ремонт завершён">')
        self.assertIn('data-f="ждём акт от асц ', html)
        self.assertIn("initStuckTable(", html)
        # комментарий можно написать прямо в таблице «Ремонты»: поле на каждый ремонт, текст текущего комментария
        self.assertRegex(html, r'<textarea class="cm" rows="2" data-et="ET0000104"[^>]*>Ждём акт от АСЦ</textarea>')
        self.assertRegex(html, r'<textarea class="cm" rows="2" data-et="ET0000103"[^>]*></textarea>')
        # письмо в АСЦ можно править: тема и текст без readonly, есть «Вернуть исходный текст»
        mail_html = self.c.get("/stuck?asc_sel=АСЦ Бета").get_data(as_text=True)
        self.assertNotIn("readonly", mail_html.split('class="mail"')[1].split("</textarea>")[0])
        self.assertIn("data-reset", mail_html)
        self.assertRegex(self.c.get("/open").get_data(as_text=True), r'<tr class="[^"]*has-cm">\s*<td class="nowrap"><a class="et" href="/repair/ET0000104"')

        def xl(q):
            wb = load_workbook(io.BytesIO(self.c.get("/stuck?export=xlsx&" + q).data))
            return sorted(r[2] for r in wb["Зависшие ремонты"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(xl("cf6=ждём"), ["ET0000104"])
        self.assertEqual(xl("cf6=(пусто)"), ["ET0000103", "ET0000103"])
        self.assertEqual(xl("cf0=принята&cf2=БЕТА"), ["ET0000103"])
        self.assertEqual(xl("cf3=казань&cf5=15"), ["ET0000103"])
        # письмо из «Листа» по выбранным ремонтам: по одному на АСЦ
        d = self.c.post("/api/request", json={"ets": ["ET0000101", "ET0000103", "ET0000104"]}).get_json()
        self.assertEqual([m["asc"] for m in d["messages"]], ["АСЦ Альфа", "АСЦ Бета"])
        self.assertEqual(d["messages"][0]["subject"], "Запрос статуса ремонта ET0000101 — АСЦ Альфа")
        self.assertIn("заявка с 01.10.2026, не передана в доставку (18 дн.)", d["messages"][0]["body"])
        self.assertIn("диагностика с 05.10.2026, результата нет (14 дн.)", d["messages"][1]["body"])
        self.assertIn('class="pick"', self.c.get("/sheet").get_data(as_text=True))

    def test_tracking_link(self):
        from datetime import datetime as D
        from testdata import repair
        rows = [repair("ET0000201", "Заявка", D(2026, 10, 1, 9), D(2026, 10, 1, 9),
                       Id="2df06f90-d104-410f-8c56-997dc88aba55"),
                repair("ET0000202", "Заявка", D(2026, 10, 1, 9), D(2026, 10, 1, 9))]
        self.run_import(make_workbook(self.tmp / "t.xlsx", rows, S1, avr=[]))
        url = "https://care.ellpa.ru/PageView/Tracking/Create/index.aspx?id=2df06f90-d104-410f-8c56-997dc88aba55"
        html = self.c.get("/repair/ET0000201").get_data(as_text=True)
        self.assertIn(f'href="{url}" target="_blank"', html)
        self.assertNotIn("Открыть в трекинге", self.c.get("/repair/ET0000202").get_data(as_text=True))
        self.assertIn(url, self.c.get("/sheet").get_data(as_text=True))
        self.assertIn(url, self.c.get("/open").get_data(as_text=True))

    def test_asc_contacts_in_emails(self):
        from datetime import datetime as D
        from openpyxl import Workbook
        from testdata import repair
        rows = [repair("ET0000301", "Заявка", D(2026, 10, 1, 9), D(2026, 10, 1, 9), asc="Альфа-Сервис",
                       КодАсцNsoВыполняющийРемонт="0304"),
                repair("ET0000302", "Диагностика", D(2026, 10, 1, 9), D(2026, 10, 2, 9), asc="Без контактов",
                       КодАсцNsoВыполняющийРемонт="9999")]
        self.run_import(make_workbook(self.tmp / "c.xlsx", rows, S1, avr=[]))
        # справочник: строка групп, заголовки, пустая строка, по строке на контактное лицо
        wb = Workbook()
        ws = wb.active
        ws.append(["Общее", None, None, None, None, None, "Контакты"])
        ws.append(["Наименование", "Статус контрагента", "Код ASC NCO", "Город", "Контактные лица", "Должность",
                   "телефон рабочий (КЛ)", "e-mail (КЛ)", "Роль контактного лица", "Банковские счета"])
        ws.append([None] * 10)
        ws.append(["Альфа-Сервис", "Активная", 304, "Москва", "Иванова Анна", "Менеджер", "+7 495 000", "Ivanova@Alfa.ru",
                   "Контактное лицо", "40702810000000000000"])
        ws.append(["Альфа-Сервис", "Активная", 304, "Москва", "Петров Пётр", "Директор", None, "boss@alfa.ru; boss2@alfa.ru",
                   "Руководитель АСЦ", "40702810000000000000"])
        ws.append(["Альфа-Сервис", "Активная", 304, "Москва", "Сидорова", None, None, None, "Логист (прямая)", None])
        path = self.tmp / "contacts.xlsx"
        wb.save(path)
        with open(path, "rb") as fh:
            r = self.c.post("/settings/contacts/upload", data={"file": (fh, "contacts.xlsx")},
                            content_type="multipart/form-data", follow_redirects=True)
        html = r.get_data(as_text=True)
        self.assertIn("Справочник загружен: 1 АСЦ, 3 адресов e-mail", html)
        self.assertIn("контакты есть у 1", html)
        raw = b""
        for fn in ("monitor.db", "monitor.db-wal"):  # банковские реквизиты в базу не попадают
            fp = self.tmp / "data" / fn
            if fp.exists():
                raw += fp.read_bytes()
        self.assertNotIn(b"40702810000000000000", raw)
        self.assertIn(b"ivanova@alfa.ru", raw)
        # письмо: «Кому» — контактное лицо (по умолчанию), остальные — для выбора
        d = self.c.post("/api/request", json={"ets": ["ET0000301", "ET0000302"]}).get_json()
        m = {x["asc"]: x for x in d["messages"]}
        self.assertEqual(m["Альфа-Сервис"]["to"], ["ivanova@alfa.ru"])
        self.assertTrue(m["Альфа-Сервис"]["mailto"].startswith("mailto:ivanova@alfa.ru?subject="))
        self.assertEqual(m["Без контактов"]["to"], [])
        self.assertIn('data-email="boss2@alfa.ru"', d["html"])
        # роли для «Копии» настраиваются
        self.c.post("/settings/contacts", data={"to_roles": ["Контактное лицо"], "cc_roles": ["Руководитель АСЦ"], "sep": ";"})
        d = self.c.post("/api/request", json={"ets": ["ET0000301"]}).get_json()
        self.assertEqual(d["messages"][0]["cc"], ["boss@alfa.ru", "boss2@alfa.ru"])
        self.assertIn("cc=boss@alfa.ru;boss2@alfa.ru", d["messages"][0]["mailto"])
        # «Где застряли» и карточка ремонта
        html = self.c.get("/stuck?asc_sel=Альфа-Сервис").get_data(as_text=True)
        self.assertIn('class="mail-to" value="ivanova@alfa.ru"', html)
        html = self.c.get("/repair/ET0000301").get_data(as_text=True)
        self.assertIn('href="mailto:ivanova@alfa.ru"', html)
        self.assertIn("Письмо в АСЦ по этому ремонту", html)
        self.assertIn("В справочнике контактов нет АСЦ с кодом 9999", self.c.get("/repair/ET0000302").get_data(as_text=True))
        # письмо в почте в браузере (Яндекс 360): ссылка на новое письмо с тем же содержимым, нужный ящик по uid
        r = self.c.post("/settings/contacts", data={"to_roles": ["Контактное лицо"], "open_in": "web", "sep": ";",
                                                     "web_url": "https://mail.360.yandex.ru/?uid=1130000066156068#/inbox"},
                        follow_redirects=True)
        self.assertIn("Настройки сохранены", r.get_data(as_text=True))
        d = self.c.post("/api/request", json={"ets": ["ET0000301"]}).get_json()
        link = d["messages"][0]["mailto"]
        self.assertTrue(link.startswith(
            "https://mail.360.yandex.ru/compose?uid=1130000066156068&mailto=mailto%3Aivanova%40alfa.ru%3Fsubject%3D"), link)
        self.assertIn('target="_blank"', d["html"])
        self.assertIn("Открыть в Яндекс Почте", d["html"])
        html = self.c.get("/repair/ET0000301").get_data(as_text=True)
        self.assertIn('href="https://mail.360.yandex.ru/compose?uid=1130000066156068&amp;mailto=mailto%3Aivanova%40alfa.ru" target="_blank"', html)
        self.assertIn("Письмо открывается в Яндекс Почте в браузере", html)
        # без адреса почты режим браузера не включается
        r = self.c.post("/settings/contacts", data={"open_in": "web", "web_url": ""}, follow_redirects=True)
        self.assertIn("вставьте адрес почты", r.get_data(as_text=True))

    def test_acts_approved_not_reported(self):
        from datetime import datetime as D
        from openpyxl import load_workbook
        from testdata import repair
        rows = [
            repair("ET0000501", "Заказаны запчасти", D(2026, 9, 23, 7), D(2026, 9, 23, 7), model="ECAM350.50.B",
                   НомерАктаНрпАто="NSE705150     ", ДатаСозданияАкта=D(2026, 9, 24, 10, 50),
                   ДатаУтвержденияАкта=D(2026, 9, 24, 15, 0), ПричинаАкта="DLG - Long delivery"),   # 25 дн. — долго
            repair("ET0000502", "Ремонт завершён", D(2026, 10, 1, 9), D(2026, 10, 15, 9), asc="АСЦ Бета",
                   НомерАктаНрпАто="ATO000123", ДатаСозданияАкта=D(2026, 10, 18, 9),
                   ДатаУтвержденияАкта=D(2026, 10, 18, 12)),                                         # 1 дн. — недавно
            repair("ET0000503", "Диагностика", D(2026, 10, 1, 9), D(2026, 10, 2, 9),
                   НомерАктаНрпАто="NSE000777", ДатаСозданияАкта=D(2026, 10, 17, 9)),                # ждёт утверждения
            repair("ET0000504", "Отчет", D(2026, 9, 1, 9), D(2026, 10, 1, 9), НомерАктаНрпАто="NSE000888",
                   ДатаСозданияАкта=D(2026, 9, 30, 9), ДатаУтвержденияАкта=D(2026, 9, 30, 10)),     # уже в «Отчет»
            repair("ET0000505", "Ремонт", D(2026, 10, 1, 9), D(2026, 10, 5, 9)),                     # акта нет
        ]
        self.run_import(make_workbook(self.tmp / "a.xlsx", rows, S1, avr=[]))
        html = self.c.get("/acts").get_data(as_text=True)
        self.assertIn('href="/acts" class="active"', html)
        ets = sorted(set(re.findall(r'href="/repair/(ET\d+)"', html)))
        self.assertEqual(ets, ["ET0000501", "ET0000502", "ET0000503"])
        kpi = dict(re.findall(r'<div class="label">([^<]+)</div><div class="value">([^<]+)</div>', html))
        self.assertEqual((kpi["Долго не в «Отчет»"], kpi["Утверждён недавно"], kpi["Акт ждёт утверждения"]), ("1", "1", "1"))
        self.assertRegex(html, r"В разделе: <b>3</b>\.\s+В выгрузке портала ремонтов с актом АНРП/АТО: 4, из них уже в «Отчет»: 1")
        for t in ("NSE705150", "АНРП", "ATO000123", "АТО", "DLG - Long delivery", "24.09.2026"):
            self.assertIn(t, html)
        # самые давние утверждения — первыми (сортировка по состоянию)
        self.assertLess(html.index("ET0000501"), html.index("ET0000502"))
        self.assertLess(html.index("ET0000502"), html.index("ET0000503"))
        # письмо в АСЦ
        d = self.c.post("/api/request", json={"ets": ["ET0000501"], "mode": "acts"}).get_json()
        m = d["messages"][0]
        self.assertEqual(m["subject"], "Перевод в «Отчет»: утверждён акт АНРП/АТО — ET0000501 — АСЦ Альфа")
        self.assertIn("— ET0000501, ECAM350.50.B: АНРП NSE705150 от 24.09.2026 утверждён 24.09.2026 (прошло 25 дн.), "
                      "ремонт в статусе «Заказаны запчасти» — просим перевести в «Отчет»", m["body"])
        # карточка ремонта: пометка и письмо про акт
        card = self.c.get("/repair/ET0000501").get_data(as_text=True)
        self.assertIn("АНРП NSE705150 от 24.09.2026", card)
        self.assertIn("Долго не в «Отчет»", card)
        self.assertIn('value="Перевод в «Отчет»: утверждён акт АНРП/АТО — ET0000501 — АСЦ Альфа"', card)
        self.assertNotIn("Долго не в «Отчет»", self.c.get("/repair/ET0000505").get_data(as_text=True))
        # порог «долго» меняется в настройках раздела
        self.c.post("/acts/settings", data={"days": "30"})
        kpi = dict(re.findall(r'<div class="label">([^<]+)</div><div class="value">([^<]+)</div>',
                              self.c.get("/acts").get_data(as_text=True)))
        self.assertEqual((kpi["Долго не в «Отчет»"], kpi["Утверждён недавно"]), ("0", "2"))
        # фильтр по состоянию и выгрузка
        html = self.c.get("/acts?ast=pending").get_data(as_text=True)
        self.assertEqual(sorted(set(re.findall(r'href="/repair/(ET\d+)"', html))), ["ET0000503"])
        wb = load_workbook(io.BytesIO(self.c.get("/acts?export=xlsx").data))
        data = list(wb.active.iter_rows(values_only=True))
        self.assertEqual(data[0][:5], ("ET", "АСЦ", "Статус", "Акт", "Акт от"))
        self.assertEqual(len(data), 4)
        self.assertEqual(sorted(r[3] for r in data[1:]), ["АНРП NSE000777", "АНРП NSE705150", "АТО ATO000123"])

    def test_models_10_days(self):
        from datetime import datetime as D
        from openpyxl import load_workbook
        from testdata import repair
        rows = [
            repair("ET0000401", "Диагностика", D(2026, 9, 30, 9), D(2026, 10, 2, 9), model="KCL95.004SI",
                   ДатаПриемаВАсц=D(2026, 10, 1), Транспортировка="Без перевозки"),           # 18 дн. — просрочен
            repair("ET0000402", "Ремонт", D(2026, 10, 10, 9), D(2026, 10, 14, 9), model="KWL90.004S",
                   ДатаПриемаВАсц=D(2026, 10, 12), Транспортировка="Клиент"),               # 7 дн. — в срок
            repair("ET0000403", "Заказаны запчасти", D(2026, 10, 8, 9), D(2026, 10, 12, 9), model="KCC9040S",
                   ДатаПриемаВАсц=D(2026, 10, 10)),                                          # 9 дн. — скоро срок
            repair("ET0000404", "Выезд/Доставка", D(2026, 10, 15, 9), D(2026, 10, 16, 9), model="CCL50.A0CP",
                   Транспортировка="Door to Door"),                                          # едет
            repair("ET0000405", "Заявка", D(2026, 10, 9, 9), D(2026, 10, 9, 9), model="KWL90.164SI",
                   Транспортировка="Клиент"),                                                # заявка, 10 дн. с оформления
            repair("ET0000406", "Ремонт завершён", D(2026, 9, 30, 9), D(2026, 10, 8, 9), model="KCL95.004SI RW",
                   ДатаПриемаВАсц=D(2026, 10, 1), ДатаЗавершенияРемонта=D(2026, 10, 8)),     # устранено за 7 дн.
            repair("ET0000407", "Диагностика", D(2026, 9, 1, 9), D(2026, 9, 2, 9), ДатаПриемаВАсц=D(2026, 9, 1)),
            repair("ET0000408", "Отчет", D(2026, 9, 1, 9), D(2026, 9, 20, 9), model="KCL95.004SI",
                   ДатаПриемаВАсц=D(2026, 9, 2)),
        ]
        self.run_import(make_workbook(self.tmp / "m.xlsx", rows, S1, avr=[]))
        html = self.c.get("/models10").get_data(as_text=True)
        ets = sorted(set(re.findall(r'href="/repair/(ET\d+)"', html)))
        self.assertEqual(ets, ["ET0000401", "ET0000402", "ET0000403", "ET0000404", "ET0000405", "ET0000406"])
        kpi = dict(re.findall(r'<div class="label">([^<]+)</div><div class="value">([^<]+)</div>', html))
        self.assertEqual(kpi["Просрочен → АНРП"], "1")
        self.assertEqual(kpi["Скоро срок"], "1")
        self.assertEqual(kpi["В срок"], "1")
        self.assertEqual(kpi["Едет к АСЦ (D2D)"], "1")
        self.assertEqual(kpi["Заявка, техника не принята"], "1")
        self.assertEqual(kpi["Завершён в срок"], "1")
        self.assertIn("+8 дн., срок был 11.10.2026", html)
        self.assertIn("едет (D2D)", html)
        # выгрузка: TAT от приёма и дни с оформления заявки
        wb = load_workbook(io.BytesIO(self.c.get("/models10?export=xlsx").data))
        data = {r[1]: r for r in wb.active.iter_rows(min_row=2, values_only=True)}
        self.assertAlmostEqual(data["ET0000402"][11], 7.46, delta=0.01)   # 12.10 00:00 → 19.10 11:00
        self.assertEqual(data["ET0000405"][10], 10)                       # 09.10 → 19.10
        self.assertIsNone(data["ET0000405"][11])
        # письмо: запрос на АНРП
        d = self.c.post("/api/request", json={"ets": ["ET0000401"], "mode": "m10"}).get_json()
        m = d["messages"][0]
        self.assertTrue(m["subject"].startswith("Запрос АНРП: превышен срок устранения дефекта (10 дн.) — ET0000401"))
        self.assertIn("Просим сформировать запрос на АНРП", m["body"])
        self.assertIn("принят в ремонт 01.10.2026, срок устранения дефекта 10 дн. истёк 11.10.2026 (просрочка 8 дн.)", m["body"])
        # карточка ремонта
        self.assertIn("Модель со сроком 10 дней от приёма в ремонт", self.c.get("/repair/ET0000401").get_data(as_text=True))
        self.assertNotIn("Модель со сроком", self.c.get("/repair/ET0000407").get_data(as_text=True))
        # настройки раздела: срок 20 дней — ET0000401 уже не просрочен
        self.c.post("/models10/settings", data={"models": "KCL95.004SI\nKWL90.004", "days": "20", "warn_days": "2", "d2d": "door"})
        html = self.c.get("/models10").get_data(as_text=True)
        kpi = dict(re.findall(r'<div class="label">([^<]+)</div><div class="value">([^<]+)</div>', html))
        self.assertEqual(kpi["Просрочен → АНРП"], "0")
        self.assertNotIn("ET0000403", html)  # KCC9040S убрали из списка моделей

    def test_repair_type_filter_persists(self):
        rows = rows_v1()[:4]
        rows[1]["ТипРемонтаПриЗавершении"] = "Дилерский клиентский"
        rows[2]["ТипРемонтаПриЗавершении"] = "Восстановление"
        rows[3]["ТипРемонтаПриЗавершении"] = None
        self.run_import(make_workbook(self.tmp / "t.xlsx", rows, S1))

        def shown(url):
            html = self.c.get(url).get_data(as_text=True)
            return sorted(set(re.findall(r'href="/repair/(ET\d+)"', html)))
        self.assertEqual(len(shown("/repairs")), 3)  # ET0000003 «Отчет» — закрыт, не показывается
        self.assertEqual(shown("/repairs?presence=closed"), ["ET0000003"])
        # выбор двух типов в фильтре
        sel = "/repairs?rtype_set=1&rtype=Клиентский+гарантийный&rtype=Дилерский+клиентский"
        self.assertEqual(shown(sel), ["ET0000001", "ET0000002"])
        # выбор запомнился: без параметров и в других разделах
        self.assertEqual(shown("/repairs"), ["ET0000001", "ET0000002"])
        html = self.c.get("/repairs").get_data(as_text=True)
        self.assertIn("Выбрано 2 из 3", html)
        self.assertEqual(shown("/attention"), [e for e in ["ET0000001", "ET0000002"] if e in shown("/attention")])
        self.assertNotIn("ET0000004", shown("/attention"))
        dash = self.c.get("/").get_data(as_text=True)
        self.assertRegex(dash, r'Ремонтов в работе</div>\s*<div class="value">2</div>')
        # «Сбросить» других фильтров тип не снимает
        self.assertEqual(shown("/repairs?status=Диагностика"), ["ET0000001"])
        # только «(пусто)»
        self.assertEqual(shown("/repairs?rtype_set=1&rtype=(пусто)"), ["ET0000004"])
        # отмечены все — фильтр снимается
        all_ = "&".join("rtype=" + v for v in ["Восстановление", "Дилерский+клиентский", "Клиентский+гарантийный", "(пусто)"])
        self.assertEqual(len(shown("/repairs?rtype_set=1&" + all_)), 3)
        self.assertIn('<span class="ms-text">Все</span>', self.c.get("/repairs").get_data(as_text=True))
        # ничего не отмечено — тоже все
        self.c.get("/repairs?rtype_set=1&rtype=Восстановление")
        self.assertEqual(shown("/repairs"), [])  # единственный такой ремонт закрыт
        self.assertEqual(shown("/repairs?presence=closed"), ["ET0000003"])
        self.assertEqual(len(shown("/repairs?rtype_set=1")), 3)
        # экспорт учитывает выбор и содержит столбец типа
        self.c.get("/repairs?rtype_set=1&rtype=(пусто)")
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(self.c.get("/repairs?export=xlsx").data))
        rows_x = list(wb.active.iter_rows(values_only=True))
        self.assertIn("Тип ремонта при завершении", rows_x[0])
        self.assertEqual(len(rows_x), 2)

    def test_contract_pages_and_settings(self):
        from datetime import datetime as dt
        from testdata import repair
        snap = dt(2026, 10, 22, 11, 0)
        rows = [
            repair("ET2000001", "Диагностика", dt(2026, 10, 12, 9), dt(2026, 10, 12, 10), ДатаПриемаВАсц=dt(2026, 10, 12),
                   ДатаДиагностикиОнЛайн=dt(2026, 10, 12, 10), TAT=10.08),
            repair("ET2000002", "Заказаны запчасти", dt(2026, 10, 5, 9), dt(2026, 10, 7, 10), ДатаПриемаВАсц=dt(2026, 10, 5)),
        ]
        res = self.run_import(make_workbook(self.tmp / "c.xlsx", rows, snap))
        self.assertEqual(res["status"], "ok", res)
        html = self.c.get("/contract").get_data(as_text=True)
        self.assertIn("Договорные сроки СЦ", html)
        self.assertIn('href="/repair/ET2000001"', html)   # просрочена диагностика
        self.assertIn("Просрочено", html)
        html = self.c.get("/contract?norm=warranty&cstate=waiting").get_data(as_text=True)
        self.assertIn('href="/repair/ET2000002"', html)
        self.assertIn("ждёт ЗЧ с 07.10.2026", html)
        # TAT из Excel как есть, договорный столбец в реестре
        html = self.c.get("/repairs?q=ET2000001").get_data(as_text=True)
        self.assertIn("10,08", html)
        self.assertIn('class="cb high"', html)
        card = self.c.get("/repair/ET2000001").get_data(as_text=True)
        self.assertIn("Договорные сроки СЦ", card)
        self.assertIn("TAT (из Excel)", card)
        # разделение в «Требуют внимания»
        self.assertIn('href="/repair/ET2000001"', self.c.get("/attention?group=contract").get_data(as_text=True))
        # экспорт договорных сроков
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(self.c.get("/contract?export=xlsx").data))
        hdr = next(wb.active.iter_rows(values_only=True))
        self.assertIn("Исключено р.д. (ожидание ЗЧ)", hdr)
        # настройка: норматив диагностики 10 р.д. -> просрочки нет
        page = self.c.get("/settings?tab=contract").get_data(as_text=True)
        form = {"enabled": "1", "accept_col": "ДатаПриемаВАсц", "country_col": "Страна",
                "type_cols": "ТипРемонтаПриЗавершении; ТипРемонта", "types": ["Клиентский гарантийный"],
                "pre_accept": ["Заявка"], "warn_days": "1", "diag_enabled": "1", "diag_days": "10",
                "diag_end_cols": ["ДатаЗаказаЗЧОнЛайн", "ДатаРемонтаОнЛайн"], "w_enabled": "1", "w_days_no_parts": "3",
                "w_days_parts": "6", "w_end_col": "ДатаЗавершенияРемонта", "w_order_col": "ДатаЗаказаЗЧОнЛайн",
                "w_delivery_col": "ДатаРемонтаОнЛайн", "w_delivery_is_proxy": "1", "w_waiting": ["Заказаны запчасти"],
                "t_enabled": "1", "t_days": "3", "t_act_col": "ДатаСозданияАкта", "t_reason_col": "ПричинаАкта"}
        self.assertIn("Договорные сроки СЦ", page)
        self.assertEqual(self.c.post("/settings/contract", data=form).status_code, 302)
        self.assertNotIn('href="/repair/ET2000001"', self.c.get("/contract?norm=diag&cstate=overdue").get_data(as_text=True))
        # а гарантийный срок просрочен при любом варианте (8 р.д. > 6) — остаётся
        self.assertIn('href="/repair/ET2000001"', self.c.get("/contract?norm=warranty&cstate=overdue").get_data(as_text=True))
        # календарь: добавить нерабочий день и проверить ошибку в дате
        cal = self.c.get("/settings?tab=calendars").get_data(as_text=True)
        self.assertIn("13.06.2025", cal)
        r = self.c.post("/settings/calendars", data={"cal_RU_present": "1", "cal_RU_name": "Россия",
                                                       "cal_RU_years": "2026", "cal_RU_holidays": "31.02.2026",
                                                       "n_countries": "0"}, follow_redirects=True)
        self.assertIn("не распознаны даты", r.get_data(as_text=True))
        r = self.c.post("/settings/calendars", data={"cal_RU_present": "1", "cal_RU_name": "Россия",
                                                       "cal_RU_years": "2026", "cal_RU_holidays": "20.10.2026 тест\n21.10.2026",
                                                       "n_countries": "1", "country_0": "РОССИЯ", "country_cal_0": "RU",
                                                       "default_calendar": "WEEKENDS", "new_code": "AM", "new_name": "Армения"},
                        follow_redirects=True)
        self.assertIn("Настройки сохранены", r.get_data(as_text=True))
        self.assertIn("Армения", self.c.get("/settings?tab=calendars").get_data(as_text=True))

    def test_export_and_picker(self):
        self.run_import(self.f1)
        r = self.c.get("/repairs?export=xlsx&presence=all")
        self.assertEqual(r.status_code, 200)
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.data))
        rows = list(wb.active.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "ET")
        self.assertEqual(len(rows), 8)
        flat = " ".join(str(v) for row in rows for v in row if v is not None)
        self.assertNotIn("Иванов", flat)
        # окно выбора файла: в среде без tkinter — понятное сообщение
        j = self.c.post("/api/pick-file").get_json()
        self.assertTrue(j["path"] or j["error"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
