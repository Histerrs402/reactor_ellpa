"""Мониторинг ремонтов — локальное веб-приложение.

Запуск:  python app.py            (откроется браузер на http://127.0.0.1:8765)
         python app.py --port 8800 --no-browser
"""
from __future__ import annotations

import argparse
import io
import json
import logging
import os
import secrets
import socket
import sys
import threading
import uuid
import webbrowser
from datetime import datetime, timedelta
from urllib.parse import urlencode

from flask import (
    Flask, abort, flash, g, jsonify, redirect, render_template, request, send_file, url_for,
)

from monitor import VERSION, analytics, charts, config, contract, db, importer, jobs, models10, openrep, portal, rules, vault
from monitor import acts as acts_mod
from monitor import tracking_api
from monitor import comments as comments_mod
from monitor import contacts as contacts_mod
from monitor.workdays import parse_date_list
from monitor.parsing import from_iso, norm_et

EVENT_LABELS = {
    "appeared": "Появился в выгрузке",
    "status_changed": "Смена статуса",
    "asc_changed": "Смена АСЦ",
    "date_changed": "Изменение даты",
    "disappeared": "Исчез из выгрузки",
    "returned": "Вернулся в выгрузку",
    "closed_by_act": "Закрытие подтверждено актом",
    "closure_manual": "Отмечен как проверенный",
    "closure_reopened": "Отметка о проверке снята",
    "excel_comment": "Комментарий из Excel",
    "excel_comment_cleared": "Комментарий в Excel очищен",
    "sheet_changed": "Перенесён в другой лист",
}
LEVELS = rules.LEVEL_LABELS
REASONS = rules.REASON_LABELS
PER_PAGE = 100

SORT_KEYS = {
    "et": lambda i: i["et"],
    "model": lambda i: (i["model"] or "").lower(),
    "status": lambda i: (i["status"] or "").lower(),
    "asc": lambda i: (i["asc"] or "").lower(),
    "rtype": lambda i: (i.get("repair_type_final") or "").lower(),
    "created": lambda i: i["created_at"],
    "status_date": lambda i: i["status_date"] or i["status_seen_since"],
    "days": lambda i: i["days_in_status"],
    "age": lambda i: i["age_days"],
    "act": lambda i: i["act_date"] or ((i["closure_info"] or {}).get("act_date") or None),
    "score": lambda i: i["score"],
    "comment": lambda i: (i["last_comment"] or {}).get("created_at"),
    "tat": lambda i: i["tat"],
    "contract": lambda i: (-contract.STATE_RANK[i["contract_worst"]] if i.get("contract_worst") else None),
    "disappeared": lambda i: i["disappeared_snapshot"],
}


def sort_items(items, key, desc):
    fn = SORT_KEYS.get(key, SORT_KEYS["score"])
    present = [i for i in items if fn(i) is not None]
    missing = [i for i in items if fn(i) is None]
    present.sort(key=fn, reverse=desc)
    return present + missing


def _secret_key() -> str:
    p = config.data_dir() / "secret.key"
    if p.exists():
        return p.read_text(encoding="utf-8").strip()
    k = secrets.token_hex(32)
    p.write_text(k, encoding="utf-8")
    return k


def _int_or_none(v):
    v = (v or "").strip()
    if v == "":
        return None
    n = int(float(v.replace(",", ".")))
    if n < 0:
        raise ValueError("отрицательное число")
    return n


def create_app() -> Flask:
    app = Flask(
        __name__,
        template_folder=str(config.resource_dir() / "templates"),
        static_folder=str(config.resource_dir() / "static"),
    )
    app.secret_key = _secret_key()
    app.config["JSON_AS_ASCII"] = False
    app.config["TEMPLATES_AUTO_RELOAD"] = False

    conn0 = db.connect()  # создаёт базу при первом запуске
    conn0.close()

    # ---------------------------------------------------------------- helpers
    def conn():
        if "conn" not in g:
            g.conn = db.connect(init=False)
        return g.conn

    @app.teardown_appcontext
    def _close(_exc):
        c = g.pop("conn", None)
        if c is not None:
            c.close()

    def settings():
        if "settings" not in g:
            g.settings = db.load_settings(conn())
        return g.settings

    def evaluation():
        if "ev" not in g:
            g.ev = rules.evaluate(conn())
        return g.ev

    @app.template_filter("d")
    def f_date(v):
        if isinstance(v, str):
            v = from_iso(v)
        return v.strftime("%d.%m.%Y") if v else "—"

    @app.template_filter("dt")
    def f_dt(v):
        if isinstance(v, str):
            v = from_iso(v)
        if not v:
            return "—"
        return v.strftime("%d.%m.%Y") if (v.hour, v.minute, v.second) == (0, 0, 0) else v.strftime("%d.%m.%Y %H:%M")

    @app.template_filter("n")
    def f_num(v):
        if v is None:
            return "—"
        try:
            return f"{int(round(v)):,}".replace(",", " ")
        except (TypeError, ValueError):
            return str(v)

    @app.template_filter("f1")
    def f_f1(v):
        if v is None:
            return "—"
        return f"{v:.1f}".replace(".", ",")

    @app.template_filter("f2")
    def f_f2(v):
        if v is None:
            return "—"
        return f"{v:.2f}".replace(".", ",")

    @app.template_filter("pct")
    def f_pct(v):
        if v is None:
            return "—"
        return f"{v * 100:.1f}".replace(".", ",") + " %"

    @app.template_filter("closure")
    def f_closure(v):
        return {"unconfirmed": "не подтверждено", "manual": "проверен вручную", "act": "закрыт актом",
                None: "не подтверждено", "": "не подтверждено"}.get(v, v)

    @app.template_filter("fromjson")
    def f_json(v):
        try:
            return json.loads(v) if v else None
        except ValueError:
            return None

    def url_with(endpoint=None, **kw):
        """URL текущей страницы (или другой) с сохранением фильтров и заменой части параметров."""
        args = request.args.to_dict(flat=False)
        for k, v in kw.items():
            if v is None:
                args.pop(k, None)
            elif isinstance(v, list):
                args[k] = v
            else:
                args[k] = [v]
        if "page" not in kw:
            args.pop("page", None)
        q = urlencode([(k, x) for k, vals in args.items() for x in vals if x != ""])
        base = url_for(endpoint) if endpoint else request.path
        return base + ("?" + q if q else "")

    def link(endpoint, **params):
        q = urlencode([(k, x) for k, v in params.items() if v not in (None, "", [])
                       for x in (v if isinstance(v, list) else [v])])
        return url_for(endpoint) + ("?" + q if q else "")

    def track_url(i):
        """Ссылка на ремонт в трекинге портала (по Id из выгрузки)."""
        from urllib.parse import quote
        tid = i.get("track_id") if isinstance(i, dict) else None
        tpl = (settings().get("tracking_url") or "").strip()
        if not tid or not tpl:
            return None
        return tpl.replace("{id}", quote(str(tid), safe="")).replace("{et}", quote(str(i.get("et") or ""), safe=""))

    @app.context_processor
    def _ctx():
        try:
            ev = evaluation()
            latest = ev["latest_import"]
            gf = {}
            for key in GLOBAL_FILTERS:  # обработчик страницы уже сохранил новый выбор, если он был
                if key in request.args and request.args.get(f"{key}_set") != "1":
                    gf[key] = request.args.getlist(key)
                elif saved_filter(key):
                    gf[key] = saved_filter(key)
            attention = sum(1 for i in analytics.apply_filters(ev["items"], {**gf, "presence": "current"}) if i["problem"]) \
                + sum(1 for i in analytics.apply_filters(ev["items"], {**gf, "presence": "unconfirmed"}) if i["problem"])
        except Exception:  # noqa: BLE001
            latest, attention, ev = None, 0, None
        as_of = ev["as_of"] if ev else None
        stale_days = (datetime.now().date() - as_of.date()).days if as_of else None
        return {
            "VERSION": VERSION,
            "APP_TITLE": config.APP_TITLE,
            "latest_import": latest,
            "as_of": as_of,
            "stale_days": stale_days,
            "attention_count": attention,
            "job": jobs.JOB.snapshot(),
            "url_with": url_with,
            "link": link,
            "track_url": track_url,
            "EVENT_LABELS": EVENT_LABELS,
            "LEVELS": LEVELS,
            "REASONS": REASONS,
            "CSTATES": contract.STATES,
            "CSTATE_CLASS": contract.STATE_CLASS,
            "CNORMS": contract.NORMS,
            "now": datetime.now(),
            "source_file": settings().get("source_file") if request else "",
            "source_mode": settings().get("source_mode") if request else "excel",
        }

    # Запоминаемые фильтры с выбором нескольких значений (как срезы в Excel):
    # ключ параметра -> поле ремонта
    GLOBAL_FILTERS = {"rtype": "repair_type_final", "status": "status", "country": "country"}

    def open_items():
        """Ремонты, из которых строятся списки значений фильтров: в текущей выгрузке, не закрытые."""
        return [i for i in evaluation()["items"] if i["current"] and not i.get("closed")]

    def ms_options(key):
        saved = db.get_meta(conn(), f"filter_{key}", []) or []
        return analytics.distinct_with_empty(open_items(), GLOBAL_FILTERS[key], extra=saved)

    def rtype_options():
        return ms_options("rtype")

    def saved_filter(key):
        """Запомненный выбор (действует во всех разделах, пока его не изменят)."""
        cache = g.setdefault("saved_filters", {})
        if key not in cache:
            cache[key] = db.get_meta(conn(), f"filter_{key}", []) or []
        return cache[key]

    def saved_rtype():
        return saved_filter("rtype")

    def get_filters():
        f = analytics.parse_filters(request.args)
        for key in GLOBAL_FILTERS:
            if request.args.get(f"{key}_set") == "1":
                sel = [v for v in request.args.getlist(key) if v]
                if set(sel) >= set(ms_options(key)):
                    sel = []  # отмечены все — значит без фильтра (новые значения тоже будут видны)
                db.set_meta(conn(), f"filter_{key}", sel)
                g.setdefault("saved_filters", {})[key] = sel
                if sel:
                    f[key] = sel
                else:
                    f.pop(key, None)
            elif key not in f and saved_filter(key):
                f[key] = list(saved_filter(key))
        return f

    def global_only(f):
        return {k: v for k, v in f.items() if k in GLOBAL_FILTERS}

    def filter_options(items, f=None):
        """Списки значений. АСЦ, бренд и модель — только среди ремонтов, подходящих под выбранные
        тип ремонта, статус и страну («подтягиваются» только подходящие)."""
        scoped = analytics.apply_filters(items, {**global_only(f or {}), "presence": "all"}) if f else items
        return {
            "rtype": ms_options("rtype"),
            "status": ms_options("status"),
            "country": ms_options("country"),
            "asc": analytics.distinct(scoped, "asc"),
            "brand": analytics.distinct(scoped, "brand"),
            "model": analytics.distinct(scoped, "model"),
            "category": analytics.distinct(scoped, "category"),
            "repair_type": analytics.distinct(scoped, "repair_type"),
            "sheet": analytics.distinct(items, "source_sheet"),
        }

    def no_data():
        return evaluation()["latest_import"] is None

    # ================================================================ главная
    @app.route("/")
    def dashboard():
        if no_data():
            return render_template("welcome.html")
        ev = evaluation()
        s = settings()
        f = get_filters()
        f.pop("presence", None)
        try:
            period = max(1, min(int(request.args.get("period") or s.get("new_period_days") or 7), 3650))
        except ValueError:
            period = 7
        items = analytics.apply_filters(ev["items"], f, "current")
        unconf = analytics.apply_filters(ev["items"], {**f, "presence": "unconfirmed"})
        as_of = ev["as_of"]
        new_since = (as_of - timedelta(days=period)) if as_of else None
        new_items = [i for i in items if i["created_at"] and new_since and i["created_at"] >= new_since]
        kpi = {
            "total": len(items),
            "problem": sum(1 for i in items if i["problem"]) + len(unconf),
            "stuck": sum(1 for i in items if i["stuck"]),
            "overdue": sum(1 for i in items if i["overdue"]),
            "date_issues": sum(1 for i in items if "date" in i["reason_kinds"]),
            "new": len(new_items),
            "unconfirmed": len(unconf),
            "detected_basis": sum(1 for i in items if i["days_basis"] == "detected"),
        }
        ckpi = {}
        controlled = [i for i in items if not i.get("demo")]
        kpi["closed"] = len(analytics.apply_filters(ev["items"], {**f, "presence": "closed"}))
        kpi["demo"] = sum(1 for i in items if i.get("demo"))
        for st in ("overdue", "soon", "waiting", "done_late", "no_data"):
            ckpi[st] = sum(1 for i in controlled if any(
                c["applicable"] and c["state"] == st for c in (i.get("contract") or {}).values()))
        ckpi_norm = {n: sum(1 for i in controlled if (i.get("contract") or {}).get(n, {}).get("state") == "overdue")
                     for n, _l in contract.NORMS}
        base = analytics.filters_to_query(f)
        base_d = {}
        for k, v in base:
            base_d.setdefault(k, []).append(v)

        def reg(**kw):
            return link("repairs", **{**base_d, **kw})

        def att(**kw):
            return link("attention", **{**base_d, **kw})

        st_rows = []
        for st, n in analytics.count_by(items, "status", "(пусто)"):
            stuck = sum(1 for i in items if (i["status"] or "(пусто)") == st and i["stuck"])
            st_rows.append({"label": st, "values": [n - stuck, stuck], "href": reg(status=st)})
        asc_counts = analytics.count_by(items, "asc", "(не указан)")
        asc_rows = []
        for a, n in asc_counts[:15]:
            pr = sum(1 for i in items if (i["asc"] or "(не указан)") == a and i["problem"])
            asc_rows.append({"label": a, "values": [n - pr, pr], "href": reg(asc=a)})
        top = sorted([i for i in items if i["problem"]] + unconf, key=lambda i: -i["score"])[:10]
        s_stuck = s["stuck"]
        s_od = s["overdue"]
        return render_template(
            "dashboard.html",
            ckpi=ckpi, ckpi_norm=ckpi_norm, contract_on=bool(s.get("contract", {}).get("enabled")), base_q=base_d,
            f=f, period=period, kpi=kpi, opts=filter_options([i for i in ev["items"] if i["current"]], f),
            status_chart=charts.hbars(st_rows, [("в пределах порога или без контроля", "c-ok"),
                                                ("дольше порога", "c-bad")]),
            asc_chart=charts.hbars(asc_rows, [("без предупреждений", "c-ok"), ("с предупреждениями", "c-warn")]),
            st_rows=st_rows, asc_rows=asc_rows, asc_total=len(asc_counts),
            top=top, reg=reg, att=att, new_since=new_since,
            stuck_enabled=s_stuck.get("enabled"), stuck_default=s_stuck.get("default_days"),
            overdue_enabled=bool(s_od.get("enabled")),
            overdue_default=s_od.get("default_days"),
            excluded_statuses=[k for k, v in s_stuck.get("per_status", {}).items() if v.get("enabled") is False],
        )

    # ========================================================== реестр
    @app.route("/repairs")
    def repairs():
        if no_data():
            return redirect(url_for("dashboard"))
        ev = evaluation()
        f = get_filters()
        items = analytics.apply_filters(ev["items"], f, "current")
        sort = request.args.get("sort", "score")
        desc = request.args.get("dir", "desc") == "desc"
        items = sort_items(items, sort, desc)
        if request.args.get("export") == "xlsx":
            return export_xlsx(items)
        try:
            page = max(1, int(request.args.get("page", 1)))
        except ValueError:
            page = 1
        pages = max(1, (len(items) + PER_PAGE - 1) // PER_PAGE)
        page = min(page, pages)
        return render_template(
            "repairs.html", f=f, items=items[(page - 1) * PER_PAGE: page * PER_PAGE], total=len(items),
            page=page, pages=pages, sort=sort, desc=desc,
            opts=filter_options(ev["items"], f), presence=f.get("presence", "current"),
        )

    def export_xlsx(items):
        from openpyxl import Workbook

        wb = Workbook(write_only=True)
        ws = wb.create_sheet("Ремонты")
        ws.append([
            "ET", "Лист", "В выгрузке", "Статус", "Дата входа в статус", "Основание расчёта дней",
            "Дней в статусе", "Порог, дн.", "АСЦ", "Код АСЦ", "Модель", "Бренд", "Категория", "Тип ремонта",
            "Дата создания", "Возраст, дн.", "Допустимый срок, дн.", "Просрочка, дн.", "Дата акта",
            "Номер акта", "Тип ремонта при завершении", "Страна", "TAT (из Excel)", "Критичность, баллы", "Уровень", "Причины",
            "Последний комментарий",
        ] + [f"{label}: состояние" for _k, label in contract.NORMS]
          + [f"{label}: прошло / допустимо, р.д." for _k, label in contract.NORMS])
        for i in items:
            ci = i["closure_info"] or {}
            ws.append([
                i["et"], i["source_sheet"],
                ("да (закрыт)" if i.get("closed") else "да") if i["current"] else ("нет (исчез)" if i["disappeared"] else "не обновлялся"),
                i["status"], i["status_date"],
                {"source": "дата из выгрузки", "detected": "от первого обнаружения (не менее)"}.get(i["days_basis"], ""),
                i["days_in_status"], i["stuck_threshold"], i["asc"], i["asc_code"], i["model"], i["brand"],
                i["category"], i["repair_type"], i["created_at"], i["age_days"], i["overdue_limit"],
                i["overdue_days"] or None, i["act_date"] or from_iso(ci.get("act_date")),
                i["act_number"] or ci.get("act_number"), i.get("repair_type_final"), i.get("country"), i["tat"], i["score"],
                "демо — не считается" if i.get("demo") else LEVELS.get(i["level"], ""),
                "; ".join(r["text"] for r in i["reasons"]),
                (i["last_comment"] or {}).get("text"),
            ] + [("не считается (демо)" if i.get("demo") else (i.get("contract") or {}).get(k, {}).get("state_label"))
                 if (i.get("contract") or {}).get(k, {}).get("applicable") else "не применяется"
                 for k, _l in contract.NORMS]
              + [(f"{c['elapsed']} / {c['norm_days']}" if c.get("elapsed") is not None and c.get("norm_days") else None)
                 for c in ((i.get("contract") or {}).get(k, {}) for k, _l in contract.NORMS)])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        name = f"ремонты_{datetime.now():%Y%m%d_%H%M}.xlsx"
        return send_file(buf, as_attachment=True, download_name=name,
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    # ========================================================== «Открытые ремонты»
    OPEN_SORT = {
        "et": lambda i: i["et"],
        "created": lambda i: i["created_at"],
        "status": lambda i: openrep.status_rank(i["status"]),
        "since": lambda i: i["status_date"] or i["status_seen_since"],
        "days": lambda i: i["days_in_status"],
        "asc": lambda i: (i["asc"] or "").lower() or None,
        "city": lambda i: (i.get("city") or "").lower() or None,
        "country": lambda i: (i.get("country") or "").lower() or None,
        "accepted": lambda i: i["accepted_at"],
        "rtype": lambda i: (i.get("repair_type_final") or "").lower() or None,
        "comment": lambda i: ((i["last_comment"] or {}).get("text") or "").lower() or None,
    }

    def sort_by(items, table, key, desc, default):
        fn = table.get(key) or table[default]
        have = [i for i in items if fn(i) is not None]
        miss = [i for i in items if fn(i) is None]
        have.sort(key=fn, reverse=desc)
        return have + miss

    @app.route("/open")
    def open_repairs():
        if no_data():
            return redirect(url_for("dashboard"))
        ev = evaluation()
        f = get_filters()
        f.pop("presence", None)
        base = analytics.apply_filters(ev["items"], f, "current")  # закрытые («Отчет») не входят
        extra = openrep.parse_extra(request.args)
        summary = openrep.status_summary(openrep.apply_extra(base, {k: v for k, v in extra.items() if k != "st"}))
        items = openrep.apply_extra(base, extra)
        sort = request.args.get("sort") or "days"
        desc = request.args.get("dir", "desc") == "desc"
        items = sort_by(items, OPEN_SORT, sort, desc, "days")
        if request.args.get("export") == "xlsx":
            return export_open(items, ev["as_of"])
        days_all = [i["days_in_status"] for i in items if i["days_in_status"] is not None]
        cities = sorted({i.get("city") or openrep.EMPTY for i in base}, key=lambda c: (c == openrep.EMPTY, c.lower()))
        return render_template(
            "open.html", f=f, extra=extra, items=items, total=len(items), summary=summary,
            n_base=len(openrep.apply_extra(base, {k: v for k, v in extra.items() if k != "st"})),
            avg_all=(sum(days_all) / len(days_all)) if days_all else None, max_all=max(days_all) if days_all else None,
            sort=sort, desc=desc, cities=cities, opts=filter_options(ev["items"], f),
            n_closed=len(analytics.apply_filters(ev["items"], {**f, "presence": "closed"})),
        )

    def export_open(items, as_of):
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Открытые ремонты"
        ws.append(["ET", "Дата создания заявки", "Статус", "Дата перехода в статус", "Дней в статусе",
                   "Дней в статусе — точность", "АСЦ", "Код АСЦ", "Город", "Страна", "Дата приёма техники",
                   "Тип ремонта", "Модель", "Последний комментарий", "Комментарий изменён"])
        for c in ws[1]:
            c.font = Font(bold=True)
        for i in items:
            lc = i["last_comment"] or {}
            ws.append([
                i["et"], i["created_at"], i["status"], i["status_date"] or i["status_seen_since"],
                i["days_in_status"],
                "не менее (даты входа в статус нет в выгрузке)" if i.get("days_lower_bound") else
                ("по дате из выгрузки" if i["days_basis"] == "source" else ""),
                i["asc"], i["asc_code"], i.get("city"), i.get("country"), i["accepted_at"],
                i.get("repair_type_final"), i["model"], lc.get("text"), from_iso(lc.get("created_at")),
            ])
            r = ws.max_row
            for col in (2, 4, 11, 15):
                ws.cell(row=r, column=col).number_format = "dd.mm.yyyy hh:mm"
        for col, w in zip("ABCDEFGHIJKLMNO", (13, 17, 18, 17, 9, 22, 32, 9, 16, 12, 17, 24, 18, 50, 17)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:O{max(ws.max_row, 1)}"
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name=f"открытые_ремонты_{datetime.now():%Y%m%d_%H%M}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    # ========================================================== «Где застряли ремонты»
    @app.route("/stuck")
    def stuck_page():
        if no_data():
            return redirect(url_for("dashboard"))
        ev = evaluation()
        s = settings()
        cfg = openrep.stuck_settings(s)
        f = get_filters()
        f.pop("presence", None)
        base = analytics.apply_filters(ev["items"], f, "current")  # закрытые («Отчет») не входят
        rows = openrep.classify(base, cfg, ev["as_of"])
        ranking = openrep.asc_ranking(rows)
        sel_cat = request.args.get("cat") or ""
        sel_asc = request.args.get("asc_sel") or ""
        shown = [r for r in rows if (not sel_cat or r["cat"] == sel_cat)
                 and (not sel_asc or (r["it"]["asc"] or "(не указан)") == sel_asc)]
        cf = openrep.parse_cf(request.args, len(openrep.STUCK_COLS))
        if request.args.get("export") == "xlsx":
            if cf:  # фильтры под заголовками таблицы «Ремонты»
                shown = [r for r in shown if openrep.cf_match(openrep.stuck_texts(r), cf)]
            return export_stuck(shown, ranking)
        statuses = openrep.status_list(base, set(s.get("closed_statuses") or []))
        in_status = {st: sum(1 for i in base if (i["status"] or openrep.EMPTY) == st
                             and not (cfg["skip_demo"] and i.get("demo"))) for st in statuses}
        stuck_n = {}
        for r in rows:
            stuck_n[r["cat"]] = stuck_n.get(r["cat"], 0) + 1
        msg = None
        if sel_asc:
            mine = [r for r in rows if (r["it"]["asc"] or "(не указан)") == sel_asc and (not sel_cat or r["cat"] == sel_cat)]
            msg = openrep.stuck_request(sel_asc, mine)
            contacts_mod.enrich(conn(), msg, mine[0]["it"]["asc_code"] if mine else None, contacts_mod.settings_cfg(s))
            msg["mailto"] = contacts_mod.mailto(msg)
        return render_template(
            "stuck.html", f=f, cfg=cfg, rows=shown, ranking=ranking, sel_cat=sel_cat, sel_asc=sel_asc, msg=msg,
            statuses=statuses, in_status=in_status, stuck_n=stuck_n, rules_of=lambda st: openrep.status_rule(cfg, st),
            cat_label=openrep.cat_label, opts=filter_options(ev["items"], f), n_ets=len({r["it"]["et"] for r in rows}),
            ACC=openrep.CAT_ACCEPTED, stuck_cols=openrep.STUCK_COLS, stuck_texts=openrep.stuck_texts,
            author=request.cookies.get("author") or s.get("author") or "",
        )

    @app.post("/stuck/settings")
    def stuck_settings_save():
        s = db.load_settings(conn())
        cfg = openrep.stuck_settings(s)
        fm = request.form
        ps = {}
        for idx in range(int(fm.get("n_st") or 0)):
            st = fm.get(f"st_{idx}")
            if not st:
                continue
            days = _int_or_none(fm.get(f"days_{idx}"))
            ps[st] = {"on": fm.get(f"on_{idx}") == "1",
                      "days": days if days is not None and 0 <= days <= 3650 else cfg["default_days"]}
        cfg["per_status"] = ps
        acc_days = _int_or_none(fm.get("acc_days"))
        cfg["accepted"] = {"on": fm.get("acc_on") == "1",
                           "days": acc_days if acc_days is not None and 0 <= acc_days <= 3650 else cfg["accepted"]["days"],
                           "statuses": fm.getlist("acc_statuses") or cfg["accepted"]["statuses"]}
        cfg["skip_demo"] = fm.get("skip_demo") == "1"
        s["stuck_tool"] = cfg
        db.save_settings(conn(), s)
        flash("Пороги сохранены.", "ok")
        return redirect(fm.get("back") or url_for("stuck_page"))

    def export_stuck(rows, ranking):
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Зависшие ремонты"
        ws.append(["АСЦ", "Код АСЦ", "ET", "Причина", "Дней", "С даты", "Статус", "Город", "Страна", "Модель",
                   "Тип ремонта", "Дата создания заявки", "Последний комментарий", "Ссылка в трекинге"])
        for r in sorted(rows, key=lambda r: ((r["it"]["asc"] or ""), r["cat"], -r["days"])):
            i = r["it"]
            ws.append([i["asc"], i["asc_code"], i["et"], openrep.cat_label(r["cat"]) + f" — {r['days']} дн.",
                       r["days"], r["since"], i["status"], i.get("city"), i.get("country"), i["model"],
                       i.get("repair_type_final"), i["created_at"], (i["last_comment"] or {}).get("text"), track_url(i)])
            for col in (6, 12):
                ws.cell(row=ws.max_row, column=col).number_format = "dd.mm.yyyy"
        w2 = wb.create_sheet("По АСЦ")
        w2.append(["АСЦ", "Код АСЦ", "Ремонтов", "По статусам", "Макс. дней", "Города"])
        for d in ranking:
            w2.append([d["asc"], d["code"], d["total"],
                       "; ".join(f"{openrep.cat_label(c)}: {n}" for c, n in d["breakdown"]), d["max_days"], d["cities"]])
        for sh, widths in ((ws, (32, 9, 13, 36, 7, 12, 18, 16, 12, 18, 24, 14, 50, 30)), (w2, (36, 9, 10, 60, 11, 40))):
            for c in sh[1]:
                c.font = Font(bold=True)
            for n, wdt in enumerate(widths, start=1):
                sh.column_dimensions[chr(64 + n)].width = wdt
            sh.freeze_panes = "A2"
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name=f"зависшие_ремонты_{datetime.now():%Y%m%d_%H%M}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    @app.post("/api/request")
    def api_request():
        """Тема и текст письма в АСЦ по выбранным ремонтам (раздел «Лист»)."""
        data = request.get_json(silent=True) or {}
        by = evaluation()["by_et"]
        items = [by[e] for e in (data.get("ets") or [])[:500] if e in by]
        if data.get("mode") == "acts":
            acfg = acts_mod.settings_cfg(settings())
            get = sheet_data()
            groups = {}
            checks = tracking_api.load_checks(conn())
            for i in items:
                r = acts_mod.evaluate(i, get(i["et"]), acfg, evaluation()["as_of"], checks.get(i["et"]))
                if r is not None:
                    groups.setdefault(i["asc"] or "(не указан)", []).append((i, r))
            msgs = [acts_mod.compose(a, groups[a]) for a in sorted(groups)]
        elif data.get("mode") == "m10":
            mcfg = models10.settings_cfg(settings())
            get = sheet_data()
            groups: dict = {}
            for i in items:
                groups.setdefault(i["asc"] or "(не указан)", []).append(
                    (i, models10.evaluate(i, get(i["et"]), mcfg, evaluation()["as_of"])))
            msgs = [models10.compose(a, groups[a], mcfg) for a in sorted(groups)]
        else:
            msgs = openrep.items_request(items)
        cfg = contacts_mod.settings_cfg(settings())
        for m in msgs:
            contacts_mod.enrich(conn(), m, by[m["ets"][0]]["asc_code"] if m["ets"] else None, cfg)
            m["mailto"] = contacts_mod.mailto(m)
        return jsonify({"ok": True, "messages": msgs, "html": render_template("_mail_list.html", msgs=msgs)})

    # ========================================================== «Лист» (как DL_GR_open_reps)
    _sheet_cache: dict = {}

    def sheet_data():
        """Исходные значения столбцов выгрузки по ET (кэш до следующего изменения данных)."""
        ver = db.data_version(conn())
        if _sheet_cache.get("ver") != ver:
            _sheet_cache.clear()
            _sheet_cache["ver"] = ver
            _sheet_cache["data"] = {r["et"]: r["data"] for r in conn().execute("SELECT et, data FROM repairs")}
            _sheet_cache["parsed"] = {}
        parsed = _sheet_cache["parsed"]

        def get(et):
            d = parsed.get(et)
            if d is None:
                raw = _sheet_cache["data"].get(et)
                d = json.loads(raw) if raw else {}
                parsed[et] = d
            return d
        return get

    def sheet_columns():
        s = settings()
        m = s["mapping"]
        special = {m.get("et"): "et", m.get("created"): "created", m.get("tat") or "TAT": "tat",
                   m.get("status"): "status", m.get("asc"): "asc"}
        cols = []
        for c in s.get("sheet_columns") or config.DEFAULT_SHEET_COLUMNS:
            kind = special.get(c, "data")
            low = c.lower()
            width = {"et": 86, "created": 104, "tat": 70}.get(kind) or next(
                (w for key, w in (("код", 48), ("бренд", 84), ("модель", 110), ("асц", 140), ("тип", 120),
                                  ("страна", 84), ("город", 92), ("регион", 104), ("результат", 120),
                                  ("статус", 110), ("дата", 104), ("сумма", 76)) if low.startswith(key) or key in low[:12]),
                104)
            cols.append({"name": c, "kind": kind, "width": width,
                         "num": kind == "tat" or low.startswith(("код", "сумма")) or c in ("TAT",)})
        return cols

    def _clean_text(s: str) -> str:
        # «_x000D_» — так Excel хранит символ возврата каретки; в тексте он не нужен
        return s.replace("_x000D_", "").replace("\r", "")

    def _fmt_cell(v):
        """(текст для показа, ключ сортировки)"""
        if v is None or v == "":
            return "", ""
        if isinstance(v, datetime):
            return (v.strftime("%d.%m.%y %H:%M") if (v.hour or v.minute) else v.strftime("%d.%m.%Y")), v.strftime("%Y-%m-%d %H:%M:%S")
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return (f"{v:.2f}".replace(".", ",") if isinstance(v, float) else str(v)), f"{float(v):020.4f}"
        sv = _clean_text(str(v))
        dt = from_iso(sv) if len(sv) >= 10 and sv[4] == "-" and sv[7] == "-" else None
        if dt:
            return _fmt_cell(dt)
        return " ".join(sv.split()), sv.lower()

    def sheet_rows(items, cols):
        get = sheet_data()
        cur = comments_mod.current_map(conn())
        rows = []
        for i in items:
            d = get(i["et"])
            cells = []
            for c in cols:
                k = c["kind"]
                if k == "et":
                    v = i["et"]
                elif k == "created":
                    v = i["created_at"]
                elif k == "tat":
                    v = i["tat"]
                elif k == "status":
                    v = i["status"]
                elif k == "asc":
                    v = i["asc"]
                else:
                    v = d.get(c["name"])
                    if c["num"] and isinstance(v, str):
                        try:
                            v = float(v.replace(",", ".")) if "." in v or "," in v else int(v)
                        except ValueError:
                            pass
                txt, key = _fmt_cell(v)
                cells.append({"text": txt, "key": key, "raw": v})
            cm = cur.get(i["et"]) or {}
            rows.append({"it": i, "cells": cells, "comment": cm.get("text") or "",
                         "comment_at": cm.get("updated_at"), "comment_by": cm.get("author")})
        return rows

    def apply_col_filters(rows, cols, args):
        flt = {}
        for idx in range(len(cols) + 1):
            v = (args.get(f"cf{idx}") or "").strip().lower()
            if v:
                flt[idx] = v
        if not flt:
            return rows
        out = []
        for r in rows:
            ok = True
            for idx, v in flt.items():
                text = r["comment"] if idx == len(cols) else r["cells"][idx]["text"]
                if v == "(пусто)":
                    if text:
                        ok = False
                        break
                elif v not in (text or "").lower():
                    ok = False
                    break
            if ok:
                out.append(r)
        return out

    @app.route("/sheet")
    def sheet():
        if no_data():
            return redirect(url_for("dashboard"))
        ev = evaluation()
        f = get_filters()
        items = analytics.apply_filters(ev["items"], f, "current")
        cst = request.args.get("cst") or ""
        if cst in ("soon", "overdue", "both"):
            want = {"soon": {"soon"}, "overdue": {"overdue"}, "both": {"soon", "overdue"}}[cst]
            items = [i for i in items if not i.get("demo") and any(
                c.get("applicable") and c.get("state") in want for c in (i.get("contract") or {}).values())]
        cols = sheet_columns()
        rows = sheet_rows(items, cols)
        cmf = request.args.get("cm") or ""
        if cmf == "with":
            rows = [r for r in rows if r["comment"]]
        elif cmf == "without":
            rows = [r for r in rows if not r["comment"]]
        tat_idx = next((n for n, c in enumerate(cols) if c["kind"] == "tat"), None)
        try:
            sidx = int(request.args.get("sort", ""))
        except ValueError:
            sidx = tat_idx if tat_idx is not None else -1
        sdesc = request.args.get("dir", "desc") == "desc"
        if 0 <= sidx < len(cols):
            filled = [r for r in rows if r["cells"][sidx]["key"] != ""]
            empty = [r for r in rows if r["cells"][sidx]["key"] == ""]
            rows = sorted(filled, key=lambda r: r["cells"][sidx]["key"], reverse=sdesc) + empty
        elif sidx == len(cols):
            rows = sorted(rows, key=lambda r: r["comment"].lower(), reverse=sdesc)
        if request.args.get("export") == "xlsx":
            return export_sheet(apply_col_filters(rows, cols, request.args), cols)
        return render_template("sheet.html", f=f, rows=rows, cols=cols, total=len(rows), cmf=cmf, cst=cst,
                               sidx=sidx, sdesc=sdesc, opts=filter_options(ev["items"], f),
                               n_comments=sum(1 for r in rows if r["comment"]),
                               author=request.cookies.get("author") or settings().get("author") or "")

    def export_sheet(rows, cols):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "Лист"
        head = [c["name"] for c in cols] + ["комментарий", "комментарий изменён"]
        ws.append(head)
        green = PatternFill("solid", fgColor="E2EFDA")
        pink = PatternFill("solid", fgColor="F8D7DA")
        hfill = PatternFill("solid", fgColor="70AD47")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = hfill
        for r in rows:
            vals = []
            for c, cell in zip(cols, r["cells"]):
                v = cell["raw"]
                if isinstance(v, str) and c["kind"] == "data":
                    dt = from_iso(v) if len(v) >= 10 and v[4] == "-" and v[7] == "-" else None
                    v = dt or _clean_text(v)
                vals.append(v)
            vals += [r["comment"] or None, from_iso(r["comment_at"]) if r["comment"] else None]
            ws.append(vals)
            row = ws.max_row
            for n, c in enumerate(cols, start=1):
                cl = ws.cell(row=row, column=n)
                if isinstance(cl.value, datetime):
                    cl.number_format = "dd.mm.yyyy hh:mm" if (cl.value.hour or cl.value.minute) else "dd.mm.yyyy"
                if c["kind"] == "tat":
                    cl.number_format = "0.00"
                    cl.fill = pink
                elif r["comment"]:
                    cl.fill = green
            ws.cell(row=row, column=len(cols) + 1).alignment = Alignment(wrap_text=False)
            if r["comment"]:
                ws.cell(row=row, column=len(cols) + 1).fill = green
            ws.cell(row=row, column=len(cols) + 2).number_format = "dd.mm.yyyy hh:mm"
        widths = {"et": 16, "created": 16, "tat": 9}
        for n, c in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(n)].width = widths.get(c["kind"], 18)
        ws.column_dimensions[get_column_letter(len(cols) + 1)].width = 50
        ws.column_dimensions[get_column_letter(len(cols) + 2)].width = 17
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{max(ws.max_row, 1)}"
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name=f"лист_{datetime.now():%Y%m%d_%H%M}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    # ========================================================== «Срок 10 дней» (отдельные модели)
    M10_COLS = [  # заголовок, ширина, вид
        ("ДатаСоздания", 104, "created"), ("ET", 104, "et"), ("Модель", 110, "model"), ("АСЦ", 140, "asc"),
        ("Код АСЦ", 50, "code"), ("Тип ремонта", 120, "rtype"), ("Город", 92, "city"), ("Статус", 110, "status"),
        ("Транспортировка", 100, "transport"), ("Дата приёма в ремонт", 96, "accepted"),
        ("Дней с оформления заявки", 84, "from_req"), ("TAT от приёма", 74, "acc_tat"), ("Срок / состояние", 236, "state"),
    ]

    def m10_rows(ev, cfg, f):
        get = sheet_data()
        cur = comments_mod.current_map(conn())
        out = []
        for i in analytics.apply_filters(ev["items"], f, "current"):
            if not models10.matches(i["model"], cfg["patterns"]):
                continue
            d = get(i["et"])
            r = models10.evaluate(i, d, cfg, ev["as_of"])
            vals = {"created": i["created_at"], "et": i["et"], "model": i["model"], "asc": i["asc"],
                    "code": i["asc_code"], "rtype": i.get("repair_type_final"), "city": i.get("city"),
                    "status": i["status"], "transport": r["transport"], "accepted": i["accepted_at"],
                    "from_req": r["from_request"], "acc_tat": r["acc_tat"]}
            cells = []
            for _name, _w, kind in M10_COLS:
                if kind == "state":
                    label, cls = models10.STATES[r["state"]]
                    det = models10.state_detail(r)
                    short = models10.state_detail(r, short=True)
                    cells.append({"text": f"{label} {det}", "key": f"{models10.severity(r):09.2f}",
                                  "html": f'<span class="badge {cls}">{label}</span>'
                                          + (f' <span class="small muted">{short}</span>' if short else ""),
                                  "title": f"{label}: {det}"})
                    continue
                txt, key = _fmt_cell(vals.get(kind))
                if kind == "acc_tat" and vals.get(kind) is not None:
                    txt = f"{vals[kind]:.2f}".replace(".", ",")
                if kind == "accepted" and not vals.get(kind):
                    txt = "едет (D2D)" if r["d2d"] else ""
                cells.append({"text": txt, "key": key, "cls": "tatv-cell" if kind == "acc_tat" else ""})
            cm = cur.get(i["et"]) or {}
            out.append({"it": i, "r": r, "cells": cells, "comment": cm.get("text") or "",
                        "comment_at": cm.get("updated_at"), "comment_by": cm.get("author")})
        return out

    @app.route("/models10")
    def models10_page():
        if no_data():
            return redirect(url_for("dashboard"))
        ev = evaluation()
        s = settings()
        cfg = models10.settings_cfg(s)
        f = get_filters()
        rows = m10_rows(ev, cfg, f)
        counts = {}
        for r in rows:
            counts[r["r"]["state"]] = counts.get(r["r"]["state"], 0) + 1
        st = request.args.get("m10") or ""
        if st in models10.STATES:
            rows = [r for r in rows if r["r"]["state"] == st]
        cmf = request.args.get("cm") or ""
        if cmf == "with":
            rows = [r for r in rows if r["comment"]]
        elif cmf == "without":
            rows = [r for r in rows if not r["comment"]]
        cols = [{"name": n, "width": w, "kind": k, "num": k in ("code", "from_req", "acc_tat"),
                 "tat_style": k == "acc_tat"} for n, w, k in M10_COLS]
        # TAT от приёма и срок закреплены справа рядом с комментарием — видны при прокрутке широкой таблицы
        right = 300
        for c in reversed(cols):
            if c["kind"] in ("state", "acc_tat"):
                c["pin_right"] = right
                right += c["width"]
        idx_tat = next(n for n, c in enumerate(cols) if c["kind"] == "acc_tat")
        try:
            sidx = int(request.args.get("sort", ""))
        except ValueError:
            sidx = idx_tat
        sdesc = request.args.get("dir", "desc") == "desc"
        if 0 <= sidx < len(cols):
            filled = [r for r in rows if r["cells"][sidx]["key"] != ""]
            empty = [r for r in rows if r["cells"][sidx]["key"] == ""]
            rows = sorted(filled, key=lambda r: r["cells"][sidx]["key"], reverse=sdesc) + empty
        if request.args.get("export") == "xlsx":
            return export_m10(apply_col_filters(rows, cols, request.args), cols)
        n = cfg["days"]
        intro = (f"Модели, для которых на устранение гарантийного дефекта дано <b>{n} дней от даты приёма в ремонт</b>; "
                 "при превышении — запрос на АНРП. TAT здесь считается от даты приёма (как TAT в Excel, но от приёма). "
                 "Даты приёма нет и доставка Door to door — техника ещё едет, срок не идёт; статус «Заявка» без D2D — "
                 "отдельный столбец «Дней с оформления заявки». Комментарии — те же, что в «Листе» (по номеру ET).")
        state_opts = [("", "Все")] + [(k, f"{models10.STATES[k][0]} ({counts.get(k, 0)})") for k in models10.STATE_ORDER]
        selects = [{"name": "m10", "label": "Состояние срока", "options": state_opts, "value": st},
                   {"name": "cm", "label": "Комментарий", "options": [("", "Все"), ("with", "С комментарием"),
                                                                     ("without", "Без комментария")], "value": cmf}]
        summary = render_template("_m10_summary.html", counts=counts, cfg=cfg, st=st, STATES=models10.STATES,
                                  ORDER=models10.STATE_ORDER)
        return render_template(
            "sheet.html", page_title=f"Срок {n} дней", page_intro=intro, action_url=url_for("models10_page"),
            extra_selects=selects, after_filters=summary, mail_mode="m10",
            f=f, rows=rows, cols=cols, total=len(rows), cmf=cmf, cst="", sidx=sidx, sdesc=sdesc,
            opts=filter_options(ev["items"], f), n_comments=sum(1 for r in rows if r["comment"]),
            author=request.cookies.get("author") or s.get("author") or "")

    @app.post("/models10/settings")
    def models10_settings():
        s = db.load_settings(conn())
        fm = request.form
        models = []
        for line in (fm.get("models") or "").replace(",", "\n").splitlines():
            m = line.strip()
            if m and m not in models:
                models.append(m)
        days = _int_or_none(fm.get("days"))
        warn = _int_or_none(fm.get("warn_days"))
        s["models10"] = {**(s.get("models10") or {}), "models": models,
                         "days": days if days and 0 < days <= 365 else 10,
                         "warn_days": warn if warn is not None and 0 <= warn <= 60 else 2,
                         "d2d": (fm.get("d2d") or "door").strip()[:40]}
        db.save_settings(conn(), s)
        flash("Настройки раздела сохранены.", "ok")
        return redirect(fm.get("back") or url_for("models10_page"))

    def export_m10(rows, cols):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "Срок 10 дней"
        ws.append([c["name"] for c in cols] + ["комментарий", "комментарий изменён"])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="70AD47")
        pink = PatternFill("solid", fgColor="F8D7DA")
        for r in rows:
            i, rr = r["it"], r["r"]
            vals = {"created": i["created_at"], "et": i["et"], "model": i["model"], "asc": i["asc"], "code": i["asc_code"],
                    "rtype": i.get("repair_type_final"), "city": i.get("city"), "status": i["status"],
                    "transport": rr["transport"], "accepted": i["accepted_at"], "from_req": rr["from_request"],
                    "acc_tat": rr["acc_tat"],
                    "state": f"{models10.STATES[rr['state']][0]}: {models10.state_detail(rr)}"}
            ws.append([vals[c["kind"]] for c in cols] + [r["comment"] or None,
                                                         from_iso(r["comment_at"]) if r["comment"] else None])
            row = ws.max_row
            for n, c in enumerate(cols, start=1):
                cl = ws.cell(row=row, column=n)
                if isinstance(cl.value, datetime):
                    cl.number_format = "dd.mm.yyyy hh:mm" if c["kind"] == "created" else "dd.mm.yyyy"
                if c["kind"] == "acc_tat":
                    cl.number_format = "0.00"
                    cl.fill = pink
            ws.cell(row=row, column=len(cols) + 2).number_format = "dd.mm.yyyy hh:mm"
        for n, c in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(n)].width = max(9, c["width"] // 7)
        ws.column_dimensions[get_column_letter(len(cols) + 1)].width = 50
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols) + 2)}{max(ws.max_row, 1)}"
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name=f"срок10_{datetime.now():%Y%m%d_%H%M}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    # ========================================================== «Акт утверждён, ремонт не в Отчете»
    ACT_COLS = [  # заголовок, ширина, вид; три последних закреплены справа рядом с комментарием
        ("ET", 104, "et"), ("АСЦ", 130, "asc"), ("Статус", 110, "status"), ("Акт", 130, "act"),
        ("Акт от", 84, "act_from"), ("Дней от даты акта", 64, "days_created"), ("Причина акта", 140, "reason"),
        ("Модель", 110, "model"), ("Город", 92, "city"), ("Код АСЦ", 50, "code"), ("ДатаСоздания", 104, "created"),
        ("Утверждён", 84, "approved"), ("Дней с утверждения", 64, "days"), ("Состояние", 110, "state"),
    ]
    ACT_PINNED = ("approved", "days", "state")

    def act_vals(i, r):
        return {"created": i["created_at"], "et": i["et"], "model": i["model"], "asc": i["asc"], "code": i["asc_code"],
                "city": i.get("city"), "status": i["status"], "act": f"{r['kind']} {r['number']}".strip(),
                "reason": r["reason"], "act_from": r["created"], "approved": r["approved"],
                "days_created": r["days_created"], "days": r["days"]}

    def act_rows(ev, cfg, f, checks):
        get = sheet_data()
        cur = comments_mod.current_map(conn())
        out = []
        for i in analytics.apply_filters(ev["items"], f, "current"):
            r = acts_mod.evaluate(i, get(i["et"]), cfg, ev["as_of"], checks.get(i["et"]))
            if r is None:
                continue
            vals = act_vals(i, r)
            cells = []
            for _name, _w, kind in ACT_COLS:
                if kind == "state":
                    label, cls = acts_mod.STATES[r["state"]]
                    det = acts_mod.state_detail(r)
                    cells.append({"text": f"{label} {det}", "key": f"{acts_mod.severity(r):09.2f}",
                                  "html": f'<span class="badge {cls}">{acts_mod.SHORT[r["state"]]}</span>',
                                  "title": f"{label}: {det}"})
                    continue
                v = vals.get(kind)
                if kind in ("act_from", "approved") and isinstance(v, datetime):
                    txt, key = v.strftime("%d.%m.%Y"), v.strftime("%Y-%m-%d %H:%M:%S")
                elif kind in ("days", "days_created") and v is not None:
                    txt, key = str(v), f"{v:06d}"
                elif kind == "approved" and not v:
                    txt, key = "ждёт", ""
                else:
                    txt, key = _fmt_cell(v)
                title = (f"утверждён {v:%d.%m.%Y %H:%M}" if kind == "approved" and isinstance(v, datetime)
                         else f"создан {v:%d.%m.%Y %H:%M}" if kind == "act_from" and isinstance(v, datetime)
                         else f"{txt} — данные: {r['source']}" + (f"; {r['note']}" if r.get("note") else "")
                         if kind == "act" else "")
                cells.append({"text": txt, "key": key, "cls": "tatv-cell" if kind == "days" else "", "title": title})
            cm = cur.get(i["et"]) or {}
            out.append({"it": i, "r": r, "cells": cells, "comment": cm.get("text") or "",
                        "comment_at": cm.get("updated_at"), "comment_by": cm.get("author")})
        return out

    @app.route("/acts")
    def acts_page():
        if no_data():
            return redirect(url_for("dashboard"))
        ev = evaluation()
        s = settings()
        cfg = acts_mod.settings_cfg(s)
        f = get_filters()
        checks = tracking_api.load_checks(conn())
        rows = act_rows(ev, cfg, f, checks)
        counts = {}
        for r in rows:
            counts[r["r"]["state"]] = counts.get(r["r"]["state"], 0) + 1
        # сколько ремонтов с актом вообще есть в данных и сколько из них уже в «Отчет» — видно, приходят ли акты
        base = analytics.apply_filters(ev["items"], {k: v for k, v in f.items() if k not in ("presence", "status")}, "all")
        with_act = [i for i in base if i["current"] and acts_mod.has_act(i)]
        diag = {"with_act": len(with_act), "closed": sum(1 for i in with_act if i.get("closed")),
                "api": sum(1 for r in rows if r["r"]["source"].startswith("API")), "checked": len(checks)}
        st = request.args.get("ast") or ""
        if st in acts_mod.STATES:
            rows = [r for r in rows if r["r"]["state"] == st]
        cmf = request.args.get("cm") or ""
        if cmf == "with":
            rows = [r for r in rows if r["comment"]]
        elif cmf == "without":
            rows = [r for r in rows if not r["comment"]]
        cols = [{"name": n, "width": w, "kind": k, "num": k in ("code", "days", "days_created"),
                 "tat_style": k == "days"} for n, w, k in ACT_COLS]
        right = 300  # дата утверждения, дни и состояние закреплены справа рядом с комментарием
        for c in reversed(cols):
            if c["kind"] in ACT_PINNED:
                c["pin_right"] = right
                right += c["width"]
        idx_state = next(n for n, c in enumerate(cols) if c["kind"] == "state")
        try:
            sidx = int(request.args.get("sort", ""))
        except ValueError:
            sidx = idx_state
        sdesc = request.args.get("dir", "desc") == "desc"
        if 0 <= sidx < len(cols):
            filled = [r for r in rows if r["cells"][sidx]["key"] != ""]
            empty = [r for r in rows if r["cells"][sidx]["key"] == ""]
            rows = sorted(filled, key=lambda r: r["cells"][sidx]["key"], reverse=sdesc) + empty
        if request.args.get("export") == "xlsx":
            return export_acts(apply_col_filters(rows, cols, request.args), cols)
        intro = ("По ремонту <b>утверждён акт АНРП (NSE…) или АТО (ATO…), а ремонт ещё не в «Отчет»</b>. "
                 "«Акт от» — дата создания акта, «Утверждён» — дата утверждения; дни — календарные на момент данных"
                 + (f" {ev['as_of']:%d.%m.%Y %H:%M}" if ev.get("as_of") else "")
                 + ". Комментарии — те же, что в «Листе» (по номеру ET).")
        state_opts = [("", "Все")] + [(k, f"{acts_mod.STATES[k][0]} ({counts.get(k, 0)})") for k in acts_mod.STATE_ORDER]
        selects = [{"name": "ast", "label": "Состояние", "options": state_opts, "value": st},
                   {"name": "cm", "label": "Комментарий", "options": [("", "Все"), ("with", "С комментарием"),
                                                                     ("without", "Без комментария")], "value": cmf}]
        tcfg = tracking_api.settings_cfg(s)
        summary = render_template("_acts_summary.html", counts=counts, cfg=cfg, st=st, STATES=acts_mod.STATES,
                                  ORDER=acts_mod.STATE_ORDER, diag=diag, n_rows=sum(counts.values()), tcfg=tcfg,
                                  tjob=tracking_api.JOB.snapshot(), tinfo=db.get_meta(conn(), "act_checks_info"),
                                  SECURE_URL=tracking_api.SECURE_URL, DOC_URL=tracking_api.DOC_URL,
                                  api_login=(s.get("api") or {}).get("login") or "",
                                  ttest=session_pop_test())
        return render_template(
            "sheet.html", page_title="Акт утверждён, ремонт не в «Отчет»", page_intro=intro,
            action_url=url_for("acts_page"), extra_selects=selects, after_filters=summary, mail_mode="acts",
            f=f, rows=rows, cols=cols, total=len(rows), cmf=cmf, cst="", sidx=sidx, sdesc=sdesc,
            opts=filter_options(ev["items"], f), n_comments=sum(1 for r in rows if r["comment"]),
            author=request.cookies.get("author") or s.get("author") or "")

    @app.post("/acts/settings")
    def acts_settings():
        s = db.load_settings(conn())
        fm = request.form
        days = _int_or_none(fm.get("days"))
        s["acts"] = {**(s.get("acts") or {}), "days": days if days is not None and 0 <= days <= 365 else 2}
        db.save_settings(conn(), s)
        flash("Настройки раздела сохранены.", "ok")
        return redirect(fm.get("back") or url_for("acts_page"))

    def export_acts(rows, cols):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "Акт утверждён"
        ws.append([c["name"] for c in cols] + ["комментарий", "комментарий изменён", "Ссылка в трекинге"])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="70AD47")
        pink = PatternFill("solid", fgColor="F8D7DA")
        green = PatternFill("solid", fgColor="E2EFDA")
        for r in rows:
            i, rr = r["it"], r["r"]
            vals = {**act_vals(i, rr), "state": f"{acts_mod.STATES[rr['state']][0]}: {acts_mod.state_detail(rr)}"}
            ws.append([vals[c["kind"]] for c in cols] + [r["comment"] or None,
                                                         from_iso(r["comment_at"]) if r["comment"] else None, track_url(i)])
            row = ws.max_row
            for n, c in enumerate(cols, start=1):
                cl = ws.cell(row=row, column=n)
                if isinstance(cl.value, datetime):
                    cl.number_format = "dd.mm.yyyy hh:mm" if c["kind"] in ("created", "approved") else "dd.mm.yyyy"
                if c["kind"] == "days":
                    cl.fill = pink
                elif r["comment"]:
                    cl.fill = green
            ws.cell(row=row, column=len(cols) + 2).number_format = "dd.mm.yyyy hh:mm"
        for n, c in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(n)].width = max(9, c["width"] // 7)
        ws.column_dimensions[get_column_letter(len(cols) + 1)].width = 50
        ws.column_dimensions[get_column_letter(len(cols) + 3)].width = 30
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols) + 3)}{max(ws.max_row, 1)}"
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name=f"акт_утверждён_{datetime.now():%Y%m%d_%H%M}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    # ---------------------------------------------------------- акты через API трекинга (SOAP)
    _ttest: dict = {}

    def session_pop_test():
        return _ttest.pop("last", None)

    @app.post("/acts/api-settings")
    def acts_api_settings():
        s = db.load_settings(conn())
        fm = request.form
        cur = {**tracking_api.DEFAULT_CFG, **(s.get("tracking_api") or {})}
        url = (fm.get("url") or "").strip() or tracking_api.SECURE_URL
        if not url.lower().startswith(("http://", "https://")):
            flash("Адрес сервиса должен начинаться с https:// (или http://).", "error")
            return redirect(fm.get("back") or url_for("acts_page"))
        workers = _int_or_none(fm.get("workers"))
        hours = _int_or_none(fm.get("recheck_hours"))
        cur.update({
            "enabled": fm.get("enabled") == "1", "auto": fm.get("auto") == "1", "url": url,
            "login": (fm.get("login") or "").strip()[:120],
            "workers": workers if workers and 1 <= workers <= 8 else 4,
            "recheck_hours": hours if hours is not None and 0 <= hours <= 24 * 30 else 20,
            "skip_statuses": [x for x in fm.getlist("skip_statuses") if x],
        })
        if fm.get("clear_password") == "1":
            cur["password"] = ""
        elif fm.get("password"):
            cur["password"] = vault.protect(fm["password"])
        s["tracking_api"] = cur
        db.save_settings(conn(), s)
        flash("Настройки проверки актов сохранены." + (
            " Внимание: адрес начинается с http:// — логин и пароль будут передаваться без шифрования."
            if url.lower().startswith("http://") else ""), "ok")
        if fm.get("then") == "test":
            return acts_api_test()
        return redirect(fm.get("back") or url_for("acts_page"))

    @app.post("/acts/api-test")
    def acts_api_test():
        fm = request.form
        cfg = tracking_api.settings_cfg(settings())
        raw = (fm.get("test_et") or "").strip()
        et = raw.upper().replace(" ", "") if tracking_api.ACT_NO.match(raw) else (norm_et(raw) or "")
        try:
            res = tracking_api.test(cfg, et or None)
            _ttest["last"] = {"ok": True, "url": cfg["url"], "insecure": cfg["insecure"], **res}
        except tracking_api.TrackingError as e:
            hint = ""
            if not cfg["insecure"]:
                hint = (f"В инструкции указан адрес {tracking_api.DOC_URL} — по нему логин и пароль передаются без "
                        f"шифрования. Можно попробовать {tracking_api.ALT_URL} (этот сервер тоже упоминается в инструкции) "
                        "или уточнить у ИТ, есть ли защищённый (https) адрес API. Подходящий адрес в браузере открывает "
                        "список методов (GetTicket, GetANRPInfo…) без страницы входа.")
            _ttest["last"] = {"ok": False, "url": cfg["url"], "error": str(e), "hint": hint, "et": et}
        return redirect((fm.get("back") or url_for("acts_page")) + "#api")

    @app.post("/acts/check/start")
    def acts_check_start():
        cfg = tracking_api.settings_cfg(settings())
        if not cfg["enabled"]:
            flash("Проверка актов через API трекинга выключена — включите её в настройках раздела.", "error")
        elif not tracking_api.JOB.start(full=request.form.get("full") == "1"):
            flash("Проверка уже идёт.", "info")
        return redirect(request.form.get("back") or url_for("acts_page"))

    @app.post("/acts/check/stop")
    def acts_check_stop():
        tracking_api.JOB.stop()
        return redirect(request.form.get("back") or url_for("acts_page"))

    @app.get("/acts/check/status")
    def acts_check_status():
        return jsonify(tracking_api.JOB.snapshot())

    @app.get("/api/ping")
    def api_ping():
        return jsonify({"app": "repair-monitor", "version": VERSION, "data_dir": str(config.data_dir())})

    @app.post("/api/comment")
    def api_comment():
        data = request.get_json(silent=True) or request.form
        et = norm_et(data.get("et")) or ""
        if not et or not conn().execute("SELECT 1 FROM repairs WHERE et=?", (et,)).fetchone():
            return jsonify({"ok": False, "error": "Ремонт не найден"}), 404
        author = (data.get("author") or request.cookies.get("author") or settings().get("author") or "").strip()[:80]
        res = comments_mod.set_current(conn(), et, data.get("text") or "", author or None)
        return jsonify({"ok": True, **res})

    @app.route("/comments/export")
    def comments_export():
        from openpyxl import Workbook
        from openpyxl.styles import Font

        c = conn()
        ev = evaluation()
        by = ev["by_et"]
        wb = Workbook()
        ws = wb.active
        ws.title = "Comments_DB"
        ws.append(["Repair", "Comment"])
        cur = [r for r in c.execute("SELECT * FROM current_comments WHERE text <> '' ORDER BY et")]
        for r in cur:
            ws.append([r["et"], r["text"]])
        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 80
        w2 = wb.create_sheet("Комментарии")
        w2.append(["ET", "Комментарий", "Изменён", "Автор", "Статус", "АСЦ", "Тип ремонта", "Страна", "В выгрузке"])
        for r in cur:
            i = by.get(r["et"]) or {}
            where = ("закрыт («" + (i.get("status") or "") + "»)" if i.get("closed") else "да") if i.get("current") \
                else ("нет (исчез)" if i.get("disappeared") else ("не обновлялся" if i else "нет в базе"))
            w2.append([r["et"], r["text"], from_iso(r["updated_at"]), r["author"], i.get("status"), i.get("asc"),
                       i.get("repair_type_final"), i.get("country"), where])
            w2.cell(row=w2.max_row, column=3).number_format = "dd.mm.yyyy hh:mm"
        for col, wd in zip("ABCDEFGHI", (14, 70, 17, 18, 18, 30, 24, 14, 16)):
            w2.column_dimensions[col].width = wd
        w3 = wb.create_sheet("История комментариев")
        w3.append(["ET", "Когда", "Автор", "Текст", "Откуда"])
        src = {"user": "программа", "excel": "Excel (Comments_DB)", "system": "действие"}
        for r in c.execute("SELECT * FROM comments ORDER BY et, id"):
            w3.append([r["et"], from_iso(r["created_at"]), r["author"], r["text"], src.get(r["source"], r["source"])])
            w3.cell(row=w3.max_row, column=2).number_format = "dd.mm.yyyy hh:mm"
        for col, wd in zip("ABCDE", (14, 17, 18, 80, 20)):
            w3.column_dimensions[col].width = wd
        for sh in (ws, w2, w3):
            for cell in sh[1]:
                cell.font = Font(bold=True)
            sh.freeze_panes = "A2"
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name=f"комментарии_{datetime.now():%Y%m%d_%H%M}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    # ========================================================== карточка
    @app.route("/repair/<et>")
    def repair(et):
        ev = evaluation()
        et = norm_et(et) or ""
        item = ev["by_et"].get(et)
        if not item:
            abort(404)
        c = conn()
        row = c.execute("SELECT data, tracked, first_seen_at FROM repairs WHERE et=?", (et,)).fetchone()
        data = json.loads(row["data"]) if row and row["data"] else {}
        tracked = json.loads(row["tracked"]) if row and row["tracked"] else {}
        events = c.execute("SELECT * FROM events WHERE et=? ORDER BY id DESC", (et,)).fetchall()
        comments = c.execute("SELECT * FROM comments WHERE et=? ORDER BY id DESC", (et,)).fetchall()
        s = settings()
        asc_contacts = contacts_mod.for_code(c, item["asc_code"])
        m10cfg = models10.settings_cfg(s)
        m10 = None
        if models10.matches(item["model"], m10cfg["patterns"]) and item["current"] and not item.get("closed"):
            m10 = models10.evaluate(item, data, m10cfg, ev["as_of"])
            m10 = {**m10, "label": models10.STATES[m10["state"]][0], "cls": models10.STATES[m10["state"]][1],
                   "detail": models10.state_detail(m10), "days": m10cfg["days"]}
        act = (acts_mod.evaluate(item, data, acts_mod.settings_cfg(s), ev["as_of"],
                                 tracking_api.load_checks(c).get(item["et"])) if item["current"] else None)
        if act:
            act = {**act, "label": acts_mod.STATES[act["state"]][0], "cls": acts_mod.STATES[act["state"]][1],
                   "title": acts_mod.act_title(act), "detail": acts_mod.state_detail(act)}
        msg = None
        if item["current"] and not item.get("closed"):
            # есть акт АНРП/АТО — письмо о переводе в «Отчет», иначе — запрос статуса
            msg = acts_mod.compose(item["asc"] or "(не указан)", [(item, act)]) if act else openrep.items_request([item])[0]
            contacts_mod.enrich(c, msg, item["asc_code"], contacts_mod.settings_cfg(s))
            msg["mailto"] = contacts_mod.mailto(msg)
        return render_template(
            "repair.html", i=item, data=data, tracked=tracked, events=events, comments=comments,
            author=request.cookies.get("author") or s.get("author") or "", token=uuid.uuid4().hex,
            status_col=s["status_date_columns"].get(item["status"] or ""), asc_contacts=asc_contacts, msg=msg,
            mcfg=contacts_mod.settings_cfg(s), email_link=contacts_mod.email_link, m10=m10, anrp=act,
            contacts_loaded=bool(db.get_meta(c, "asc_contacts_info")),
        )

    def _add_comment(et, text, author, token, source="user"):
        return comments_mod.add(conn(), et, text, author, token, source)

    @app.post("/repair/<et>/comment")
    def repair_comment(et):
        et = norm_et(et) or ""
        if not conn().execute("SELECT 1 FROM repairs WHERE et=?", (et,)).fetchone():
            abort(404)
        text = (request.form.get("text") or "").strip()
        author = (request.form.get("author") or "").strip()[:80]
        if not text:
            flash("Комментарий пустой — ничего не сохранено.", "warn")
        elif _add_comment(et, text[:4000], author, request.form.get("token")):
            flash("Комментарий сохранён.", "ok")
        resp = redirect(url_for("repair", et=et) + "#comments")
        if author:
            resp.set_cookie("author", author, max_age=3600 * 24 * 365)
        return resp

    @app.post("/repair/<et>/resolve")
    def repair_resolve(et):
        et = norm_et(et) or ""
        c = conn()
        row = c.execute("SELECT present, closure_status FROM repairs WHERE et=?", (et,)).fetchone()
        if not row:
            abort(404)
        action = request.form.get("action", "resolve")
        note = (request.form.get("note") or "").strip()[:1000]
        author = (request.form.get("author") or "").strip()[:80]
        now = db.now_iso()
        if action == "resolve" and not row["present"]:
            info = {"by": author, "note": note, "at": now}
            c.execute("UPDATE repairs SET closure_status='manual', closure_info=? WHERE et=?",
                      (json.dumps(info, ensure_ascii=False), et))
            c.execute(
                "INSERT INTO events(import_id, et, detected_at, event_type, field, old_value, new_value, note) "
                "VALUES(NULL,?,?,'closure_manual','',?, 'manual', ?)",
                (et, now, row["closure_status"], (f"{author}: " if author else "") + (note or "без пояснения")),
            )
            _add_comment(et, "Отмечено как проверенное (исчез из выгрузки)" + (f": {note}" if note else ""),
                         author, None, source="system")
            db.bump_version(c)
            flash("Ремонт отмечен как проверенный и убран из списка «Требуют внимания».", "ok")
        elif action == "reopen" and not row["present"] and row["closure_status"] == "manual":
            c.execute("UPDATE repairs SET closure_status='unconfirmed', closure_info=NULL WHERE et=?", (et,))
            c.execute(
                "INSERT INTO events(import_id, et, detected_at, event_type, field, old_value, new_value, note) "
                "VALUES(NULL,?,?,'closure_reopened','', 'manual', 'unconfirmed', ?)",
                (et, now, author or None),
            )
            db.bump_version(c)
            flash("Отметка снята — ремонт снова требует проверки.", "ok")
        return redirect(request.form.get("back") or url_for("repair", et=et))

    # ===================================================== требуют внимания
    @app.route("/attention")
    def attention():
        if no_data():
            return redirect(url_for("dashboard"))
        ev = evaluation()
        f = get_filters()
        f.pop("presence", None)
        cur = analytics.apply_filters(ev["items"], f, "current")
        dis = analytics.apply_filters(ev["items"], {**f, "presence": "unconfirmed"})
        items = [i for i in cur if i["problem"]] + dis
        counts = {k: sum(1 for i in items if k in i["reason_kinds"]) for k in REASONS}
        group = request.args.get("group", "")
        groups = {
            "contract": sum(1 for i in items if "contract" in i["reason_kinds"]),
            "internal": sum(1 for i in items if set(i["reason_kinds"]) - {"contract"}),
            "": len(items),
        }
        if group == "contract":
            items = [i for i in items if "contract" in i["reason_kinds"]]
        elif group == "internal":
            items = [i for i in items if set(i["reason_kinds"]) - {"contract"}]
        if "reason" in f or "level" in f:
            items = [i for i in items
                     if (not f.get("reason") or set(f["reason"]) & set(i["reason_kinds"]))
                     and (not f.get("level") or i["level"] in f["level"])]
        sort = request.args.get("sort", "score")
        desc = request.args.get("dir", "desc") == "desc"
        items = sort_items(items, sort, desc)
        levels ={k: sum(1 for i in items if i["level"] == k) for k in LEVELS}
        try:
            page = max(1, int(request.args.get("page", 1)))
        except ValueError:
            page = 1
        pages = max(1, (len(items) + PER_PAGE - 1) // PER_PAGE)
        page = min(page, pages)
        shown = items[(page - 1) * PER_PAGE: page * PER_PAGE]
        hist: dict = {}
        if shown:
            ets = [i["et"] for i in shown]
            q = (f"SELECT * FROM events WHERE et IN ({','.join('?' * len(ets))}) "
                 "AND NOT (event_type='appeared' AND note='первичная загрузка листа') ORDER BY id DESC")
            for e in conn().execute(q, ets):
                lst = hist.setdefault(e["et"], [])
                if len(lst) < 4:
                    lst.append(e)
        return render_template(
            "attention.html", f=f, items=shown, total=len(items), page=page, pages=pages,
            counts=counts, levels=levels, hist=hist, sort=sort, desc=desc, group=group, groups=groups,
            opts=filter_options([i for i in ev["items"] if i["current"] or i["disappeared"]], f),
            author=request.cookies.get("author") or settings().get("author") or "",
        )

    # ======================================================== договорные сроки
    DEFAULT_CSTATES = ["overdue", "soon", "done_late", "waiting", "no_data"]

    def contract_rows(items, sel_norm, sel_states):
        rows = []
        for i in items:
            for n, _l in contract.NORMS:
                if sel_norm and n != sel_norm:
                    continue
                c = (i.get("contract") or {}).get(n)
                if not c or not c["applicable"] or c["state"] == "na":
                    continue
                if c["state"] not in sel_states:
                    continue
                rows.append((i, c))
        rows.sort(key=lambda r: (contract.STATE_RANK[r[1]["state"]], -(r[1]["overdue_by"] or 0),
                                 r[1]["remaining"] if r[1]["remaining"] is not None else 999, r[0]["et"]))
        return rows

    @app.route("/contract")
    def contract_page():
        if no_data():
            return redirect(url_for("dashboard"))
        ev = evaluation()
        s = settings()
        f = get_filters()
        f.pop("presence", None)
        items = [i for i in analytics.apply_filters(ev["items"], f, "current") if not i.get("demo")]
        n_demo = sum(1 for i in analytics.apply_filters(ev["items"], f, "current") if i.get("demo"))
        sel_norm = request.args.get("norm") or ""
        if sel_norm not in dict(contract.NORMS):
            sel_norm = ""
        sel_states = [x for x in request.args.getlist("cstate") if x in contract.STATES] or DEFAULT_CSTATES
        matrix = {n: {st: 0 for st in contract.STATES} for n, _l in contract.NORMS}
        for i in items:
            for n, _l in contract.NORMS:
                c = (i.get("contract") or {}).get(n)
                if c:
                    matrix[n][c["state"]] += 1
        rows = contract_rows(items, sel_norm, sel_states)
        if request.args.get("export") == "xlsx":
            return export_contract(rows)
        # соблюдение по АСЦ: завершённые операции и текущие просрочки
        small = int(s.get("small_sample") or 10)
        by_asc: dict = {}
        for i in items:
            a = i["asc"] or "(не указан)"
            d = by_asc.setdefault(a, {"asc": a, "n": 0, **{n: {"done": 0, "late": 0, "overdue": 0} for n, _l in contract.NORMS}})
            d["n"] += 1
            for n, _l in contract.NORMS:
                c = (i.get("contract") or {}).get(n)
                if not c or not c["applicable"]:
                    continue
                if c["state"] in ("done_ok", "done_late"):
                    d[n]["done"] += 1
                if c["state"] == "done_late":
                    d[n]["late"] += 1
                if c["state"] == "overdue":
                    d[n]["overdue"] += 1
        asc_rows = list(by_asc.values())
        for d in asc_rows:
            d["violations"] = sum(d[n]["late"] + d[n]["overdue"] for n, _l in contract.NORMS)
            d["done_total"] = sum(d[n]["done"] for n, _l in contract.NORMS)
        asc_rows.sort(key=lambda d: (-d["violations"], -d["n"]))
        try:
            page = max(1, int(request.args.get("page", 1)))
        except ValueError:
            page = 1
        per = 200
        pages = max(1, (len(rows) + per - 1) // per)
        page = min(page, pages)
        return render_template(
            "contract.html", f=f, matrix=matrix, rows=rows[(page - 1) * per: page * per], total=len(rows),
            page=page, pages=pages, sel_norm=sel_norm, sel_states=sel_states, asc_rows=asc_rows[:60],
            asc_total=len(asc_rows), small=small, cs=s["contract"], n_items=len(items), n_demo=n_demo,
            opts=filter_options([i for i in ev["items"] if i["current"]], f),
        )

    def export_contract(rows):
        from openpyxl import Workbook

        wb = Workbook(write_only=True)
        ws = wb.create_sheet("Договорные сроки")
        ws.append(["ET", "АСЦ", "Страна", "Статус", "Тип ремонта", "Норматив", "Вариант", "Допустимо, р.д.",
                   "Начало отсчёта (приём)", "Контрольная дата", "Контрольная операция", "Прошло р.д.",
                   "Исключено р.д. (ожидание ЗЧ)", "Осталось р.д.", "Просрочка р.д.", "Состояние",
                   "Приближение", "Основание", "Недостающие данные", "Календарь"])
        for i, c in rows:
            ws.append([
                i["et"], i["asc"], i["country"], i["status"], i.get("repair_type_final"), c["label"],
                {"A": "без ЗЧ", "B": "с ЗЧ"}.get(c["variant"], ""), c["norm_days"], c["start"], c["control"],
                c["control_label"], c["elapsed"], c["excluded"] or None, c["remaining"], c["overdue_by"],
                c["state_label"], "да" if c["approx"] else "", "; ".join(c["basis"]), "; ".join(c["missing"]),
                c["calendar"],
            ])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name=f"договорные_сроки_{datetime.now():%Y%m%d_%H%M}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    # ============================================================ аналитика
    @app.route("/analytics")
    def analytics_page():
        if no_data():
            return redirect(url_for("dashboard"))
        ev = evaluation()
        s = settings()
        f = get_filters()
        f.pop("presence", None)
        items = analytics.apply_filters(ev["items"], f, "current")
        st = analytics.status_table(items, s)
        asc = analytics.asc_table(items, s)
        asc_sort = request.args.get("asc_sort", "problem")
        keyf = {
            "problem": lambda r: (-r["problem"], -r["n"]),
            "problem_share": lambda r: (r["small"], -r["problem_share"], -r["n"]),
            "stuck": lambda r: (-r["stuck"], -r["n"]),
            "stuck_share": lambda r: (r["small"], -r["stuck_share"], -r["n"]),
            "n": lambda r: -r["n"],
            "median": lambda r: -(r["median_days"] or -1),
            "name": lambda r: r["asc"].lower(),
        }.get(asc_sort, lambda r: (-r["problem"], -r["n"]))
        asc["rows"].sort(key=keyf)
        asc_all = request.args.get("asc_all") == "1"
        asc["n_rows"] = len(asc["rows"])
        if not asc_all:
            asc["rows"] = asc["rows"][:40]
        tr = analytics.transitions(conn(), int(s.get("small_sample") or 10))
        dyn = analytics.dynamics(conn())
        dyn_chart = charts.line_chart(dyn, [
            ("problem", "С предупреждениями (в выгрузке)", "c-s1"),
            ("stuck", "Дольше порога", "c-s2"),
            ("overdue", "Просрочен срок", "c-s3"),
            ("disappeared_unconfirmed", "Исчезли, не подтв.", "c-s4"),
        ])
        stuck_rows = [{"label": r["status"], "values": [r["stuck"]],
                       "href": link("attention", status=r["status"], reason="stuck")}
                      for r in st if r["stuck"]]
        return render_template(
            "analytics.html", f=f, st=st, asc=asc, tr=tr, dyn=dyn, dyn_chart=dyn_chart,
            stuck_chart=charts.hbars(stuck_rows, [("дольше порога", "c-bad")]),
            asc_sort=asc_sort, asc_all=asc_all, opts=filter_options([i for i in ev["items"] if i["current"]], f),
            n_items=len(items), small=int(s.get("small_sample") or 10),
            transitions_total=sum(r["n_all"] for r in tr),
        )

    # ================================================================ история
    @app.route("/history")
    def history():
        c = conn()
        et = (request.args.get("q") or "").strip().upper()
        etype = request.args.get("type") or ""
        d_from = request.args.get("from") or ""
        d_to = request.args.get("to") or ""
        show_initial = request.args.get("initial") == "1"
        where, params = [], []
        if et:
            where.append("et LIKE ?")
            params.append(f"%{et}%")
        if etype:
            where.append("event_type=?")
            params.append(etype)
        if d_from:
            where.append("detected_at >= ?")
            params.append(d_from)
        if d_to:
            where.append("detected_at < date(?, '+1 day')")
            params.append(d_to)
        if not show_initial:
            where.append("NOT (event_type='appeared' AND note='первичная загрузка листа')")
        w = ("WHERE " + " AND ".join(where)) if where else ""
        total = c.execute(f"SELECT COUNT(*) FROM events {w}", params).fetchone()[0]
        try:
            page = max(1, int(request.args.get("page", 1)))
        except ValueError:
            page = 1
        per = 200
        pages = max(1, (total + per - 1) // per)
        page = min(page, pages)
        events = c.execute(f"SELECT * FROM events {w} ORDER BY id DESC LIMIT ? OFFSET ?",
                           params + [per, (page - 1) * per]).fetchall()
        type_counts = c.execute("SELECT event_type, COUNT(*) n FROM events GROUP BY event_type ORDER BY n DESC").fetchall()
        imports = c.execute("SELECT * FROM imports ORDER BY id DESC LIMIT 30").fetchall()
        return render_template("history.html", events=events, total=total, page=page, pages=pages,
                               type_counts=type_counts, imports=imports, q=et, etype=etype,
                               d_from=d_from, d_to=d_to, show_initial=show_initial)

    # ============================================================ комментарии
    @app.route("/comments", methods=["GET", "POST"])
    def comments():
        c = conn()
        if request.method == "POST":
            et = norm_et(request.form.get("et")) or ""
            text = (request.form.get("text") or "").strip()
            author = (request.form.get("author") or "").strip()[:80]
            if not et or not c.execute("SELECT 1 FROM repairs WHERE et=?", (et,)).fetchone():
                flash(f"Ремонт «{et or '—'}» не найден. Проверьте номер ET.", "error")
            elif not text:
                flash("Комментарий пустой — ничего не сохранено.", "warn")
            elif _add_comment(et, text[:4000], author, request.form.get("token")):
                flash(f"Комментарий к {et} сохранён.", "ok")
            resp = redirect(url_for("comments"))
            if author:
                resp.set_cookie("author", author, max_age=3600 * 24 * 365)
            return resp
        q = (request.args.get("q") or "").strip()
        src = request.args.get("source") or ""
        author = (request.args.get("author") or "").strip()
        where, params = [], []
        if q:
            where.append("(et LIKE ? OR text LIKE ?)")
            params += [f"%{q.upper()}%", f"%{q}%"]
        if src:
            where.append("source=?")
            params.append(src)
        if author:
            where.append("author LIKE ?")
            params.append(f"%{author}%")
        w = ("WHERE " + " AND ".join(where)) if where else ""
        total = c.execute(f"SELECT COUNT(*) FROM comments {w}", params).fetchone()[0]
        try:
            page = max(1, int(request.args.get("page", 1)))
        except ValueError:
            page = 1
        per = 200
        pages = max(1, (total + per - 1) // per)
        page = min(page, pages)
        rows = c.execute(f"SELECT * FROM comments {w} ORDER BY id DESC LIMIT ? OFFSET ?",
                         params + [per, (page - 1) * per]).fetchall()
        by_src = {r["source"]: r["n"] for r in c.execute("SELECT source, COUNT(*) n FROM comments GROUP BY source")}
        ev = evaluation()
        return render_template("comments.html", rows=rows, total=total, page=page, pages=pages, q=q,
                               src=src, author_q=author, by_src=by_src, by_et=ev["by_et"],
                               author=request.cookies.get("author") or settings().get("author") or "",
                               token=uuid.uuid4().hex, prefill=request.args.get("et", ""))

    # ============================================================== загрузка
    @app.route("/import")
    def import_page():
        s = settings()
        last = conn().execute("SELECT * FROM imports ORDER BY id DESC LIMIT 1").fetchone()
        return render_template("import.html", s=s, last=last, check=None, AUTH_LABELS=portal.AUTH_LABELS,
                               path=request.args.get("path") or s.get("source_file") or "")

    @app.post("/import/check")
    def import_check():
        s = settings()
        path = (request.form.get("path") or "").strip().strip('"')
        try:
            rep = importer.check_file(path, s)
            db.set_meta(conn(), "workbook_info", {"file": path, "sheets": rep["sheets"], "headers": rep["headers"]})
        except importer.UserError as e:
            rep = {"ok": False, "errors": [str(e)], "warnings": [], "infos": [], "sheets": [], "per_sheet": {}}
        except Exception as e:  # noqa: BLE001
            rep = {"ok": False, "errors": [f"Не удалось проверить файл: {e}"], "warnings": [], "infos": [],
                   "sheets": [], "per_sheet": {}}
        last = conn().execute("SELECT * FROM imports ORDER BY id DESC LIMIT 1").fetchone()
        return render_template("import.html", s=s, last=last, check=rep, path=path, AUTH_LABELS=portal.AUTH_LABELS)

    @app.post("/import/start")
    def import_start():
        force = request.form.get("force") == "1"
        if request.form.get("mode") == "api":
            if not jobs.JOB.start("", force=force, mode="api"):
                flash("Загрузка уже выполняется — дождитесь её окончания.", "warn")
            return redirect(url_for("import_page"))
        path = (request.form.get("path") or settings().get("source_file") or "").strip().strip('"')
        if not path:
            flash("Сначала укажите файл выгрузки.", "warn")
            return redirect(url_for("import_page"))
        force = request.form.get("force") == "1"
        if not jobs.JOB.start(path, force=force):
            flash("Загрузка уже выполняется — дождитесь её окончания.", "warn")
        return redirect(url_for("import_page", path=path))

    @app.post("/import/api-test")
    def import_api_test():
        s = settings()
        expected = [s["mapping"].get(k) for k in config.REQUIRED_FIELDS] + [s["mapping"].get("created")]
        try:
            res = portal.test_connection(s["api"], [c for c in expected if c])
        except importer.UserError as e:
            res = {"ok": False, "error": str(e)}
        except Exception as e:  # noqa: BLE001
            res = {"ok": False, "error": f"{e.__class__.__name__}: {e}"}
        last = conn().execute("SELECT * FROM imports ORDER BY id DESC LIMIT 1").fetchone()
        return render_template("import.html", s=s, last=last, check=None, api_check=res,
                               path=s.get("source_file") or "", AUTH_LABELS=portal.AUTH_LABELS)

    @app.get("/import/status")
    def import_status():
        st = jobs.JOB.snapshot()
        st["elapsed"] = round(__import__("time").time() - st["started_at"], 1) if st.get("started_at") else None
        return jsonify(st)

    @app.post("/api/pick-file")
    def pick_file():
        path, err = jobs.pick_file_subprocess(request.form.get("initial") or settings().get("source_file") or "")
        return jsonify({"path": path, "error": err})

    # ============================================================== настройки
    def known_lists():
        c = conn()
        s = settings()
        wbi = db.get_meta(c, "workbook_info", {}) or {}
        headers = []
        for sh in s.get("sheets", []):
            for h in (wbi.get("headers") or {}).get(sh, []):
                if h not in headers:
                    headers.append(h)
        statuses = [r["status"] for r in c.execute(
            "SELECT status, COUNT(*) n FROM repairs WHERE status IS NOT NULL GROUP BY status ORDER BY n DESC")]
        for st in list(s["status_date_columns"]) + list(s["stuck"]["per_status"]) + list(s["overdue"]["per_status"]):
            if st and st not in statuses:
                statuses.append(st)
        ascs = [r["asc"] for r in c.execute(
            "SELECT asc, COUNT(*) n FROM repairs WHERE present=1 AND asc IS NOT NULL GROUP BY asc ORDER BY asc")]
        cat_field = s["overdue"].get("category_field") or ""
        ev = evaluation()
        cats = sorted({(i.get("overdue_category") or "") for i in ev["items"] if i["current"]} - {""},
                      key=str.lower)
        for k in s["overdue"]["per_category"]:
            if k not in cats:
                cats.append(k)
        cs = s["contract"]

        def data_distinct(col):
            vals = set()
            for i in ev["items"]:
                d = i.get("_data")
                if d and d.get(col):
                    vals.add(str(d[col]).strip())
            return sorted(v for v in vals if v)

        reason_col = cs["tech"].get("reason_col") or ""
        act_reasons = data_distinct(reason_col) if reason_col else []
        for r in cs["tech"].get("exclude_reasons") or []:
            if r not in act_reasons:
                act_reasons.append(r)
        type_vals = []
        for col in cs.get("type_cols") or []:
            for v in data_distinct(col):
                if v not in type_vals:
                    type_vals.append(v)
        for v in cs.get("types") or []:
            if v not in type_vals:
                type_vals.append(v)
        countries = sorted({(i.get("country") or "").strip().upper() for i in ev["items"]} - {""}
                           | set(cs.get("country_map") or {}))
        if not headers:
            headers = sorted({h for i in ev["items"] for h in (i.get("_data") or {})})
        return {"headers": headers, "date_headers": [h for h in headers if "дата" in h.lower()], "statuses": statuses,
                "ascs": ascs, "cats": cats, "cat_field": cat_field,
                "sheets": wbi.get("sheets") or [], "all_headers": wbi.get("headers") or {}, "wb_file": wbi.get("file"),
                "act_reasons": sorted(act_reasons, key=str.lower), "type_vals": sorted(type_vals, key=str.lower),
                "countries": countries, "part_headers": [h for h in headers if "партномер" in h.lower()]}

    @app.route("/settings")
    def settings_page():
        s = settings()
        tab = request.args.get("tab", "source")
        cinfo = None
        if tab == "contacts":
            c = conn()
            codes = contacts_mod.all_codes(c)
            asc_rows = c.execute("SELECT asc_code, asc, COUNT(*) n FROM repairs WHERE present=1 GROUP BY asc_code, asc "
                                 "ORDER BY n DESC").fetchall()
            missing = [dict(r) for r in asc_rows if contacts_mod.norm_code(r["asc_code"]) not in codes]
            cinfo = {"meta": db.get_meta(c, "asc_contacts_info"), "codes": len(codes),
                     "contacts": c.execute("SELECT COUNT(*) FROM asc_contacts").fetchone()[0],
                     "with_email": c.execute("SELECT COUNT(*) FROM asc_contacts WHERE email IS NOT NULL").fetchone()[0],
                     "asc_total": len(asc_rows), "missing": missing,
                     "roles": sorted(set(contacts_mod.roles(c)) | set(contacts_mod.settings_cfg(s)["to_roles"])
                                     | set(contacts_mod.settings_cfg(s)["cc_roles"])),
                     "cfg": contacts_mod.settings_cfg(s)}
        return render_template("settings.html", s=s, k=known_lists(), FIELDS=config.FIELDS,
                               DATE_CHECKS=config.DATE_CHECKS, OPS=rules.OPS, backups=db.list_backups(),
                               data_dir=str(config.data_dir()), tab=tab, cinfo=cinfo,
                               AUTH_LABELS=portal.AUTH_LABELS, pw_protected=vault.is_protected(s["api"].get("password")))

    @app.post("/settings/<section>")
    def settings_save(section):
        s = db.load_settings(conn())
        fm = request.form
        try:
            if section == "source":
                if fm.get("source_mode") in ("excel", "api"):
                    s["source_mode"] = fm.get("source_mode")
                s["source_file"] = (fm.get("source_file") or "").strip().strip('"')
                s["sheets"] = fm.getlist("sheets") or [x for x in (fm.get("sheets_text") or "").split(";") if x.strip()]
                s["closure_sheet"] = fm.get("closure_sheet", "")
                s["comments_sheet"] = fm.get("comments_sheet", "")
            elif section == "api":
                api = s["api"]
                if fm.get("source_mode") in ("excel", "api"):
                    s["source_mode"] = fm.get("source_mode")
                api["base_url"] = (fm.get("base_url") or "").strip() or portal.DEFAULT_BASE_URL
                api["auth"] = fm.get("auth") if fm.get("auth") in portal.AUTH_LABELS else "basic"
                api["login"] = (fm.get("login") or "").strip()
                if fm.get("password"):
                    api["password"] = vault.protect(fm.get("password"))
                elif fm.get("clear_password") == "1":
                    api["password"] = ""
                api["token_header"] = (fm.get("token_header") or "").strip() or "X-API-Key"
                if fm.get("token"):
                    api["token"] = vault.protect(fm.get("token").strip())
                elif fm.get("clear_token") == "1":
                    api["token"] = ""
                api["timeout"] = max(10, min(_int_or_none(fm.get("timeout")) or 120, 1800))
                api["verify_ssl"] = fm.get("verify_ssl") == "1"
                api["check_closures"] = fm.get("check_closures") == "1"
                try:
                    api["time_shift_hours"] = float((fm.get("time_shift_hours") or "0").replace(",", "."))
                except ValueError:
                    raise ValueError("поправка времени должна быть числом часов") from None
                srcs = []
                for idx in range(int(fm.get("n_src") or 0) + 1):
                    name = (fm.get(f"src_name_{idx}") or "").strip()
                    if not name:
                        continue
                    srcs.append({"name": name, "enabled": fm.get(f"src_on_{idx}") == "1",
                                 "filter": " ".join((fm.get(f"src_filter_{idx}") or "").split())})
                if not srcs:
                    raise ValueError("нужен хотя бы один запрос к порталу")
                if len({x["name"] for x in srcs}) != len(srcs):
                    raise ValueError("названия запросов не должны повторяться")
                api["sources"] = srcs
            elif section == "contacts":
                web_url = (fm.get("web_url") or "").strip()[:500]
                open_in = "web" if fm.get("open_in") == "web" else "mailto"
                if open_in == "web" and not contacts_mod.web_base(web_url):
                    raise ValueError("для почты в браузере вставьте адрес почты из строки браузера")
                s["asc_contacts"] = {"to_roles": fm.getlist("to_roles"), "cc_roles": fm.getlist("cc_roles"),
                                     "sep": fm.get("sep") if fm.get("sep") in (";", ",") else ";",
                                     "open_in": open_in, "web_url": web_url}
            elif section == "sheet":
                cols = []
                for line in (fm.get("sheet_columns") or "").splitlines():
                    c = line.strip()
                    if c and c not in cols and c.lower() != "комментарий":
                        cols.append(c)
                s["sheet_columns"] = cols or list(config.DEFAULT_SHEET_COLUMNS)
            elif section == "mapping":
                for key, *_ in config.FIELDS:
                    if f"map_{key}" in fm:
                        s["mapping"][key] = fm.get(f"map_{key}") or ""
                for key in ("et", "text"):
                    s["comments_mapping"][key] = (fm.get(f"cm_{key}") or "").strip()
                for key in ("et", "act_date", "act_number", "status"):
                    s["closure_mapping"][key] = (fm.get(f"zm_{key}") or "").strip()
                missing = [config.FIELD_LABELS[k] for k in config.REQUIRED_FIELDS if not s["mapping"].get(k)]
                if missing:
                    raise ValueError("не заданы обязательные поля: " + ", ".join(missing))
            elif section == "statuses":
                sdc, sps, ops = {}, {}, {}
                n = int(fm.get("n_status") or 0)
                for idx in range(n + 1):
                    name = (fm.get(f"st_{idx}") or "").strip()
                    if not name:
                        continue
                    col = fm.get(f"sd_{idx}") or ""
                    if col:
                        sdc[name] = col
                    sp = {"enabled": fm.get(f"stuck_on_{idx}") == "1", "days": _int_or_none(fm.get(f"stuck_days_{idx}"))}
                    if not sp["enabled"] or sp["days"] is not None:
                        sps[name] = sp
                    op = {"enabled": fm.get(f"od_on_{idx}") == "1", "days": _int_or_none(fm.get(f"od_days_{idx}"))}
                    if not op["enabled"] or op["days"] is not None:
                        ops[name] = op
                s["status_date_columns"] = sdc
                s["stuck"]["per_status"] = sps
                s["overdue"]["per_status"] = ops
                s["early_statuses"] = fm.getlist("early")
                s["closed_statuses"] = fm.getlist("closed")
            elif section == "stuck":
                s["stuck"]["enabled"] = fm.get("enabled") == "1"
                s["stuck"]["default_days"] = _int_or_none(fm.get("default_days"))
                pa, psa = {}, []
                for idx in range(int(fm.get("n_asc") or 0) + 1):
                    a = (fm.get(f"asc_{idx}") or "").strip()
                    d = _int_or_none(fm.get(f"asc_days_{idx}"))
                    if a and d is not None:
                        pa[a] = {"days": d}
                for idx in range(int(fm.get("n_sa") or 0) + 1):
                    st = (fm.get(f"sa_status_{idx}") or "").strip()
                    a = (fm.get(f"sa_asc_{idx}") or "").strip()
                    d = _int_or_none(fm.get(f"sa_days_{idx}"))
                    if st and a and d is not None:
                        psa.append({"status": st, "asc": a, "days": d})
                s["stuck"]["per_asc"] = pa
                s["stuck"]["per_status_asc"] = psa
            elif section == "overdue":
                s["overdue"]["enabled"] = fm.get("enabled") == "1"
                s["overdue"]["default_days"] = _int_or_none(fm.get("default_days"))
                new_field = fm.get("category_field") or ""
                if new_field != s["overdue"].get("category_field"):
                    s["overdue"]["category_field"] = new_field
                    s["overdue"]["per_category"] = {}
                else:
                    pc = {}
                    for idx in range(int(fm.get("n_cat") or 0) + 1):
                        name = (fm.get(f"cat_{idx}") or "").strip()
                        if not name:
                            continue
                        on = fm.get(f"cat_on_{idx}") == "1"
                        d = _int_or_none(fm.get(f"cat_days_{idx}"))
                        if not on or d is not None:
                            pc[name] = {"enabled": on, "days": d}
                    s["overdue"]["per_category"] = pc
            elif section == "checks":
                for code, _l in config.DATE_CHECKS:
                    s["date_checks"][code] = fm.get(f"chk_{code}") == "1"
            elif section == "rules":
                out = []
                for idx in range(int(fm.get("n_rules") or 0) + 1):
                    fld = (fm.get(f"r_field_{idx}") or "").strip()
                    if not fld or fm.get(f"r_del_{idx}") == "1":
                        continue
                    op = fm.get(f"r_op_{idx}") or "eq"
                    if op not in rules.OPS:
                        op = "eq"
                    out.append({
                        "id": fm.get(f"r_id_{idx}") or uuid.uuid4().hex[:8],
                        "enabled": fm.get(f"r_on_{idx}") == "1",
                        "label": (fm.get(f"r_label_{idx}") or fld).strip()[:80],
                        "field": fld, "op": op, "value": (fm.get(f"r_value_{idx}") or "").strip(),
                        "points": _int_or_none(fm.get(f"r_points_{idx}")) or 0,
                    })
                s["custom_rules"] = out
            elif section == "weights":
                for k in s["weights"]:
                    v = _int_or_none(fm.get(f"w_{k}"))
                    s["weights"][k] = v if v is not None else 0
                for k in ("high", "medium"):
                    v = _int_or_none(fm.get(f"l_{k}"))
                    if v is not None:
                        s["levels"][k] = v
                s["small_sample"] = _int_or_none(fm.get("small_sample")) or 10
                s["new_period_days"] = _int_or_none(fm.get("new_period_days")) or 7
            elif section == "history":
                s["tracked_date_fields"] = fm.getlist("tracked")
                s["excluded_columns"] = [x.strip() for x in (fm.get("excluded") or "").splitlines() if x.strip()]
            elif section == "contract":
                cs = s["contract"]
                cs["enabled"] = fm.get("enabled") == "1"
                for key in ("accept_col", "country_col"):
                    cs[key] = fm.get(key) or cs[key]
                cs["type_cols"] = [x.strip() for x in (fm.get("type_cols") or "").split(";") if x.strip()]
                cs["types"] = fm.getlist("types")
                cs["pre_accept_statuses"] = fm.getlist("pre_accept")
                cs["include_start_day"] = fm.get("include_start_day") == "1"
                cs["exclude_order_day"] = fm.get("exclude_order_day") == "1"
                cs["exclude_delivery_day"] = fm.get("exclude_delivery_day") == "1"
                cs["warn_days"] = _int_or_none(fm.get("warn_days")) or 0
                cs["attention_done_late"] = fm.get("attention_done_late") == "1"
                cs["attention_no_data"] = fm.get("attention_no_data") == "1"
                dg = cs["diag"]
                dg["enabled"] = fm.get("diag_enabled") == "1"
                dg["days"] = _int_or_none(fm.get("diag_days")) or 3
                dg["end_cols"] = fm.getlist("diag_end_cols")
                w = cs["warranty"]
                w["enabled"] = fm.get("w_enabled") == "1"
                w["days_no_parts"] = _int_or_none(fm.get("w_days_no_parts")) or 3
                w["days_parts"] = _int_or_none(fm.get("w_days_parts")) or 6
                for key in ("end_col", "order_col", "delivery_col", "parts_sum_col"):
                    w[key] = fm.get("w_" + key) or ""
                w["delivery_is_proxy"] = fm.get("w_delivery_is_proxy") == "1"
                w["parts_cols"] = fm.getlist("w_parts_cols")
                w["waiting_statuses"] = fm.getlist("w_waiting")
                w["skip_if_tech"] = fm.get("w_skip_if_tech") == "1"
                t = cs["tech"]
                t["enabled"] = fm.get("t_enabled") == "1"
                t["days"] = _int_or_none(fm.get("t_days")) or 3
                for key in ("act_col", "act_number_col", "reason_col", "pending_field"):
                    t[key] = fm.get("t_" + key) or ""
                t["exclude_reasons"] = fm.getlist("t_exclude")
                t["pending_values"] = [x.strip() for x in (fm.get("t_pending_values") or "").split(";") if x.strip()]
                if not cs["types"]:
                    raise ValueError("не выбран ни один тип ремонта под договорным контролем")
            elif section == "calendars":
                cals = s["calendars"]
                errors = []
                for code in list(cals):
                    if fm.get(f"cal_{code}_present") != "1":
                        continue
                    c = cals[code]
                    c["name"] = (fm.get(f"cal_{code}_name") or code).strip()[:60]
                    c["source"] = (fm.get(f"cal_{code}_source") or "").strip()[:300]
                    try:
                        c["years"] = sorted({int(y) for y in (fm.get(f"cal_{code}_years") or "").replace(",", " ").split()})
                    except ValueError:
                        raise ValueError(f"годы календаря «{c['name']}» — укажите числами через запятую")
                    hol, e1 = parse_date_list((fm.get(f"cal_{code}_holidays") or "").splitlines())
                    wd, e2 = parse_date_list((fm.get(f"cal_{code}_workdays") or "").splitlines())
                    errors += e1 + e2
                    c["holidays"], c["workdays"] = hol, wd
                if errors:
                    raise ValueError("не распознаны даты: " + ", ".join(errors[:5]))
                new_code = (fm.get("new_code") or "").strip().upper()
                if new_code:
                    if not new_code.isalnum() or len(new_code) > 12:
                        raise ValueError("код календаря — латинские буквы и цифры, до 12 символов")
                    cals.setdefault(new_code, {"name": (fm.get("new_name") or new_code).strip(), "years": [],
                                               "holidays": [], "workdays": [], "source": ""})
                cmap = {}
                for idx in range(int(fm.get("n_countries") or 0)):
                    country = (fm.get(f"country_{idx}") or "").strip().upper()
                    code = fm.get(f"country_cal_{idx}") or ""
                    if country and code in cals:
                        cmap[country] = code
                s["contract"]["country_map"] = cmap
                dc = fm.get("default_calendar")
                if dc in cals:
                    s["contract"]["default_calendar"] = dc
            elif section == "general":
                s["author"] = (fm.get("author") or "").strip()[:80]
                s["backup_keep"] = _int_or_none(fm.get("backup_keep")) or 15
                if "tracking_url" in fm:
                    s["tracking_url"] = (fm.get("tracking_url") or "").strip()[:500]
                if "demo_keywords" in fm:
                    s["demo_keywords"] = [k.strip() for k in (fm.get("demo_keywords") or "").split(";") if k.strip()][:20]
            else:
                abort(404)
        except ValueError as e:
            flash(f"Настройки не сохранены: {e}.", "error")
            return redirect(url_for("settings_page", tab="source" if section == "api" else section))
        db.save_settings(conn(), s)
        flash("Настройки сохранены. Расчёты пересчитаны по новым правилам.", "ok")
        return redirect(url_for("settings_page", tab="source" if section == "api" else section))

    @app.post("/settings/contacts/upload")
    def contacts_upload():
        f = request.files.get("file")
        if not f or not f.filename:
            flash("Выберите файл со справочником контрагентов (Excel).", "warn")
            return redirect(url_for("settings_page", tab="contacts"))
        tmpdir = config.data_dir() / "tmp"
        tmpdir.mkdir(parents=True, exist_ok=True)
        src = tmpdir / f"contacts_{uuid.uuid4().hex}.xlsx"
        try:
            f.save(str(src))
            rows, info = contacts_mod.parse(str(src), tmpdir)
            info["file"] = f.filename
            contacts_mod.save(conn(), rows, info)
            n_codes = len({r["code"] for r in rows})
            flash(f"Справочник загружен: {n_codes} АСЦ, {sum(1 for r in rows if r['email'])} адресов e-mail.", "ok")
        except importer.UserError as e:
            flash(f"Справочник не загружен: {e}", "error")
        except Exception as e:  # noqa: BLE001
            flash(f"Справочник не загружен: {e.__class__.__name__}: {e}", "error")
        finally:
            try:
                src.unlink()
            except OSError:
                pass
        return redirect(url_for("settings_page", tab="contacts"))

    @app.post("/backup")
    def make_backup():
        p = db.backup(conn(), "manual", keep=None)
        flash(f"Резервная копия создана: {p.name}", "ok")
        return redirect(url_for("settings_page", tab="general"))

    @app.errorhandler(404)
    def not_found(_e):
        return render_template("error.html", code=404, message="Страница или ремонт не найдены."), 404

    @app.errorhandler(500)
    def server_error(e):
        logging.getLogger("monitor").exception("500")
        importer._log_error(repr(getattr(e, "original_exception", e)))
        return render_template("error.html", code=500,
                               message="Внутренняя ошибка. Подробности записаны в data/logs/errors.log."), 500

    return app


# ====================================================================== запуск
def _free_port(start: int) -> int:
    for p in range(start, start + 50):
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sck:
            if os.name != "nt":  # порт после недавнего закрытия программы (TIME_WAIT) свободен
                sck.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
            try:
                sck.bind(("127.0.0.1", p))
                return p
            except OSError:
                continue
    return start


def _running_instance(port: int) -> str | None:
    """Адрес уже запущенной программы с этой же папкой данных (чтобы не запускать вторую копию)."""
    import urllib.request
    url = f"http://127.0.0.1:{port}/"
    try:
        opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
        with opener.open(url + "api/ping", timeout=1.5) as r:
            info = json.loads(r.read().decode("utf-8"))
    except Exception:  # noqa: BLE001
        return None
    if info.get("app") == "repair-monitor" and info.get("data_dir") == str(config.data_dir()):
        return url
    return None


def main(argv=None):
    ap = argparse.ArgumentParser(description="Мониторинг ремонтов")
    ap.add_argument("--port", type=int, default=8765)
    ap.add_argument("--no-browser", action="store_true")
    ap.add_argument("--pick-file", nargs="?", const="", default=None, help=argparse.SUPPRESS)
    args = ap.parse_args(argv)
    if args.pick_file is not None:
        sys.exit(jobs.run_file_dialog(args.pick_file))

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        handlers=[logging.FileHandler(config.data_dir() / "logs" / "app.log", encoding="utf-8")],
    )
    logging.getLogger("werkzeug").setLevel(logging.WARNING)
    running = _running_instance(args.port)
    if running:
        print(f"Программа уже запущена: {running} — открываю её в браузере.")
        if not args.no_browser:
            webbrowser.open(running)
        return
    app = create_app()
    port = _free_port(args.port)
    url = f"http://127.0.0.1:{port}/"
    print(f"{config.APP_TITLE} {VERSION}")
    print(f"Открыто в браузере: {url}")
    print(f"Данные программы: {config.data_dir()}")
    print("Чтобы остановить программу, закройте это окно.")
    if not args.no_browser:
        threading.Timer(1.2, lambda: webbrowser.open(url)).start()
    os.environ.setdefault("FLASK_RUN_FROM_CLI", "false")
    app.run(host="127.0.0.1", port=port, debug=False, use_reloader=False, threaded=True)


if __name__ == "__main__":
    main()
